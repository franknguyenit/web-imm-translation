"""Cửa 0 của bộ công cụ: chạy `python3 -m unittest discover -s tests` từ gốc dự án."""
import csv, importlib.util, json, shutil, subprocess, sys, tempfile, unittest
from pathlib import Path

GOC = Path(__file__).resolve().parent.parent
MAU = GOC / "tests" / "mau" / "viec-mau"
spec = importlib.util.spec_from_file_location("dich", GOC / "cong-cu" / "dich.py")
dich = importlib.util.module_from_spec(spec); spec.loader.exec_module(dich)


def chay_kiem(viec):
    r = subprocess.run([sys.executable, str(GOC / "cong-cu" / "dich.py"), "kiem", str(viec)], capture_output=True, text=True)
    return r.returncode, (viec / "kiem-bao-cao.md").read_text(encoding="utf-8")


class KiemCua0(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp()); self.viec = self.tmp / "viec"
        shutil.copytree(MAU, self.viec)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def sua(self, id_, cot, gia_tri):
        p = self.viec / "song-ngu.tsv"
        rows = dich.doc_tsv(p)
        for r in rows:
            if r["id"] == id_:
                r[cot] = gia_tri
        dich.ghi_tsv(p, rows)

    def sua_seo(self, k, v):
        p = self.viec / "ban-giao" / "seo.json"; d = json.loads(p.read_text()); d[k] = v; p.write_text(json.dumps(d))

    def do(self, mau_loi):
        ma, bc = chay_kiem(self.viec)
        self.assertEqual(ma, 1, bc)
        chan = bc.split("## CẢNH BÁO")[0]
        self.assertIn(mau_loi, chan, bc)

    def test_bai_mau_xanh(self):
        ma, bc = chay_kiem(self.viec)
        self.assertEqual(ma, 0, bc)

    def test_so_lech(self):
        self.sua("s002", "en", "The minimum investment is **US$900,000**, and Form I-526E petitions must be filed by September 30, 2026."); self.do("SỐ LỆCH")

    def test_ngay_lech(self):
        self.sua("s002", "en", "The minimum investment is **US$800,000**, and Form I-526E petitions must be filed by September 30, 2025."); self.do("SỐ LỆCH")

    def test_quy_doi_sai(self):
        self.sua("s004", "en", "Service fees are approximately US$35,000."); self.do("QUY ĐỔI SAI")

    def test_con_vnd(self):
        self.sua("s004", "en", "Service fees are approximately 520 million VND (US$20,000)."); self.do("CÒN ĐƠN VỊ VND")

    def test_thuat_ngu_chot(self):
        self.sua("s003", "en", "Your whole family receives a temporary green card, including your spouse and unmarried children under 21"); self.do("THUẬT NGỮ «thẻ xanh có điều kiện»")

    def test_khong_dich(self):
        self.sua("s007", "en", "More than 2,500 clients have trusted our firm"); self.do("THUẬT NGỮ «IMM Group»")

    def test_con_tieng_viet(self):
        self.sua("s007", "en", "More than 2,500 khách hàng have trusted IMM Group"); self.do("CÒN CHỮ TIẾNG VIỆT")

    def test_tu_cam(self):
        self.sua("s003", "en", "Your whole family receives a guaranteed green card: a conditional green card for your spouse and children under 21"); self.do("TỪ CẤM")

    def test_khach_viet(self):
        self.sua("s007", "en", "More than 2,500 Vietnamese clients have trusted IMM Group"); self.do("TỪ CẤM")

    def test_thieu_ban_dich(self):
        self.sua("s006", "en", ""); self.do("THIẾU BẢN DỊCH")

    def test_bo_khong_ly_do(self):
        self.sua("s008", "ghi_chu", ""); self.do("ĐÃ BỎ đoạn")

    def test_anh_thieu_alt(self):
        self.sua("s005", "en", ""); self.do("ẢNH THIẾU ALT")

    def test_mat_lien_ket(self):
        self.sua("s006", "en", "Learn more about the EB-5 program."); self.do("SỐ LIÊN KẾT LỆCH")

    def test_hai_h1(self):
        self.sua("s007", "loai", "h1"); self.do("cần đúng 1 H1")

    def test_doan_trich_khong_h1(self):
        self.sua("s001", "loai", "h2")
        self.do("cần đúng 1 H1")
        self.sua_seo("phan_doan", True)
        ma, bc = chay_kiem(self.viec)
        self.assertEqual(ma, 0, bc)
        self.assertIn("đoạn trích", bc)
        self.sua("s007", "loai", "h1"); self.sua("s001", "loai", "h1")
        self.do("cần đúng 1 H1")  # cờ đoạn trích không tha 2 H1

    def test_seo_title_dai(self):
        self.sua_seo("meta_title", "EB-5 Visa: The Complete Guide to a U.S. Green Card Through Investment for Families"); self.do("meta title")

    def test_seo_mo_ta_ngan(self):
        self.sua_seo("meta_description", "EB-5 visa guide."); self.do("meta description")

    def test_slug_sai(self):
        self.sua_seo("slug", "EB5 Visa_Mỹ"); self.do("slug")

    def test_tu_khoa_khong_trong_title(self):
        self.sua_seo("meta_title", "U.S. Green Card Through Investment | IMM Group"); self.do("không có trong meta title")

    def test_chua_ty_gia(self):
        (self.viec / "ty-gia.json").unlink(); self.do("chưa chạy `tygia`")


class DonVi(unittest.TestCase):
    def test_so_vi(self):
        self.assertEqual(dich.lay_so("800.000 USD, 2,5%, 1 tỷ đồng, 30/09/2026", "vi"), ([800000.0, 2.5, 30.0, 9.0, 2026.0], [1e9]))

    def test_so_en(self):
        self.assertEqual(sorted(dich.lay_so("US$800,000, 2.5%, three years, September 30th, 2026, $1.2 million", "en")[0]),
                         sorted([800000.0, 2.5, 3.0, 9.0, 30.0, 2026.0, 1200000.0]))

    def test_may_khong_phai_thang(self):
        self.assertEqual(dich.lay_so("Rules may change. You may save. Filed May 22, 2026 or 1 March 2027.", "en")[0], [5.0, 22.0, 2026.0, 1.0, 3.0, 2027.0])
        self.assertEqual(dich.lay_so("EB-5 may not pay. Rules may change in 20 years.", "en")[0], [5.0, 20.0])

    def test_thang_nam(self):
        self.assertEqual(dich.lay_so("Founded in January 2005; June 2016 | VTC8", "en")[0], [1.0, 2005.0, 6.0, 2016.0, 8.0])

    def test_ten_nuoc_ngoai_khong_phai_tieng_viet(self):
        self.assertFalse(dich.con_tieng_viet("André Miranda of Estudio Jurídico"))
        self.assertTrue(dich.con_tieng_viet("more than khách hàng"))
        self.assertTrue(dich.con_tieng_viet("Phục vụ"))

    def test_tach_md(self):
        d = dich.tach_md("# Tiêu đề\n\nĐoạn một\ndòng hai\n\n- mục a\n![ảnh](a.png)\n> trích")
        self.assertEqual([x["loai"] for x in d], ["h1", "p", "li", "img", "quote"])
        self.assertEqual(d[1]["vi"], "Đoạn một dòng hai")

    def test_tach_html_bo_rac(self):
        m, d = dich.tach_html('<html><head><title>T</title><link rel="alternate" hreflang="en" href="https://x/en/"></head><body><div id="content"><!-- CTA --><h1>Tiêu đề <span class="material-symbols-outlined">info</span></h1><p>Xin <a href="/a">chào</a></p><footer>chân</footer><img src="b.webp" alt="ảnh"></div></body></html>')
        self.assertEqual(m["hreflang_en"], "https://x/en/")
        self.assertEqual([(x["loai"], x["vi"]) for x in d], [("h1", "Tiêu đề"), ("p", "Xin [chào](/a)"), ("img", "ảnh")])

    def test_nap_xoa_cu_giu_moi(self):
        tmp = Path(tempfile.mkdtemp())
        cu_tm, cu_tmx = dich.TM, dich.TMX
        try:
            dich.TM, dich.TMX = tmp / "tm.jsonl", tmp / "tm.tmx"
            dich.TM.write_text(json.dumps({"vi": "Thẻ xanh", "en": "Green card cũ", "loai": "p", "viec": "a", "nguon": "", "ngay": ""}, ensure_ascii=False) + "\n")
            v = tmp / "v"; v.mkdir()
            dich.ghi_tsv(v / "song-ngu.tsv", [{"id": "s1", "loai": "p", "vi": "thẻ  xanh", "en": "Green card"}, {"id": "s2", "loai": "p", "vi": "Mới", "en": "New"}])
            dich.lenh_nap([str(v)])
            tm = dich.doc_tm()
            self.assertEqual([e["en"] for e in tm], ["Green card", "New"])
            self.assertEqual(dich.goi_y_tm("Thẻ xanh", tm), "100%: Green card")
            self.assertIn("<tmx", dich.TMX.read_text())
        finally:
            dich.TM, dich.TMX = cu_tm, cu_tmx; shutil.rmtree(tmp)

    def test_dien(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            dich.ghi_tsv(tmp / "song-ngu.tsv", [{"id": "s001", "loai": "p", "vi": "a"}, {"id": "s002", "loai": "p", "vi": "b", "ghi_chu": "trung:s001"}])
            (tmp / "d.txt").write_text("[s001] Hello\nworld\n[s002|#bo: chỉ cho khách Việt] [BO]\n")
            dich.lenh_dien([str(tmp), str(tmp / "d.txt")])
            r = dich.doc_tsv(tmp / "song-ngu.tsv")
            self.assertEqual(r[0]["en"], "Hello world")
            self.assertEqual((r[1]["en"], r[1]["ghi_chu"]), ("[BO]", "trung:s001 #bo: chỉ cho khách Việt"))
            dich.ghi_tsv(tmp / "song-ngu.tsv", r + [{"id": "s003", "loai": "p", "vi": "c", "tm": "100%: C"}, {"id": "s004", "loai": "p", "vi": "d", "tm": "90%: D"}])
            dich.lenh_dien([str(tmp), "--tm100"])
            r2 = dich.doc_tsv(tmp / "song-ngu.tsv")
            self.assertEqual([x["en"] for x in r2], ["Hello world", "[BO]", "C", ""])
            dich.ghi_tsv(tmp / "song-ngu.tsv", [{"id": "s001", "loai": "p", "vi": "a"}, {"id": "s002", "loai": "p", "vi": "a", "ghi_chu": "trung:s001"}])
            (tmp / "t.txt").write_text("[s001] A")
            dich.lenh_dien([str(tmp), str(tmp / "t.txt"), "--trung"])
            self.assertEqual([x["en"] for x in dich.doc_tsv(tmp / "song-ngu.tsv")], ["A", "A"])
            (tmp / "x.txt").write_text("[s999] lạc")
            with self.assertRaises(SystemExit):
                dich.lenh_dien([str(tmp), str(tmp / "x.txt")])
        finally:
            shutil.rmtree(tmp)

    def test_ghep_giu_so_thu_tu(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            dich.ghi_tsv(tmp / "song-ngu.tsv", [{"id": "s001", "loai": "li", "vi": "5. a", "en": "5. A?"}, {"id": "s002", "loai": "li", "vi": "b", "en": "B"}])
            dich.lenh_ghep([str(tmp)])
            self.assertEqual((tmp / "ban-giao" / "bai-dich.en.md").read_text().splitlines()[:2], ["5. A?", "- B"])
        finally:
            shutil.rmtree(tmp)

    def test_ghep_ra_ban_vi_thang_hang(self):
        """bai-dich.vi.md phải cùng số dòng, cùng cấp heading với bai-dich.en.md — kể cả khi có đoạn bỏ."""
        tmp = Path(tempfile.mkdtemp())
        try:
            dich.ghi_tsv(tmp / "song-ngu.tsv", [
                {"id": "s001", "loai": "tieu-de", "vi": "Tiêu đề trang", "en": "Page title"},
                {"id": "s002", "loai": "h2", "vi": "Điều kiện {1}EB-5{/1}", "en": "{1}EB-5{/1} requirements"},
                {"id": "s003", "loai": "li", "vi": "Vốn 800.000 USD", "en": "US$800,000 capital"},
                {"id": "s004", "loai": "li", "vi": "Tạo 10 việc làm", "en": "Creates 10 jobs"},
                {"id": "s005", "loai": "p", "vi": "Gọi văn phòng TP.HCM", "en": "[BO]", "ghi_chu": "#bo: chỉ hợp khách Việt"},
                {"id": "s006", "loai": "p", "vi": "Liên hệ tư vấn", "en": "Talk to an advisor"},
            ])
            dich.lenh_ghep([str(tmp)])
            en = (tmp / "ban-giao" / "bai-dich.en.md").read_text().splitlines()
            vi = (tmp / "ban-giao" / "bai-dich.vi.md").read_text().splitlines()
            self.assertEqual(len(en), len(vi))
            self.assertEqual(en, ["## {1}EB-5{/1} requirements".replace("{1}", "").replace("{/1}", ""),
                                  "", "- US$800,000 capital", "- Creates 10 jobs", "", "Talk to an advisor"])
            self.assertEqual(vi, ["## Điều kiện EB-5", "", "- Vốn 800.000 USD", "- Tạo 10 việc làm", "", "Liên hệ tư vấn"])
        finally:
            shutil.rmtree(tmp)

    def test_ghep_chi_vi_khong_dung_ban_en(self):
        """--chi-vi lấp bản VI cho việc cũ mà không ghi lại bản EN đã qua cửa 0."""
        tmp = Path(tempfile.mkdtemp())
        try:
            dich.ghi_tsv(tmp / "song-ngu.tsv", [{"id": "s001", "loai": "p", "vi": "xin chào", "en": "hello"}])
            (tmp / "ban-giao").mkdir()
            (tmp / "ban-giao" / "bai-dich.en.md").write_text("đã sửa tay, không được đụng\n")
            dich.lenh_ghep([str(tmp), "--chi-vi"])
            self.assertEqual((tmp / "ban-giao" / "bai-dich.en.md").read_text(), "đã sửa tay, không được đụng\n")
            self.assertEqual((tmp / "ban-giao" / "bai-dich.vi.md").read_text(), "xin chào\n")
            self.assertFalse((tmp / "ban-giao" / "bai-dich.en.html").exists())
        finally:
            shutil.rmtree(tmp)

    def test_trung_khong_lan_tu_alt_anh(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            dich.ghi_tsv(tmp / "song-ngu.tsv", [{"id": "s001", "loai": "img", "vi": "IBDA", "en": "IBDA logo"}, {"id": "s002", "loai": "p", "vi": "IBDA", "ghi_chu": "trung:s001"}])
            dich.lenh_dien([str(tmp), "--trung"])
            self.assertEqual(dich.doc_tsv(tmp / "song-ngu.tsv")[1]["en"], "")
        finally:
            shutil.rmtree(tmp)

    def test_bang_thuat_ngu_that_xanh(self):
        self.assertEqual(dich.kiem_bang_thuat_ngu(dich.doc_thuat_ngu()), [])

    def test_tm_tach_alt_anh_voi_chu(self):
        tmp = Path(tempfile.mkdtemp())
        cu_tm, cu_tmx = dich.TM, dich.TMX
        try:
            dich.TM, dich.TMX = tmp / "tm.jsonl", tmp / "tm.tmx"
            v = tmp / "v"; v.mkdir()
            dich.ghi_tsv(v / "song-ngu.tsv", [{"id": "s001", "loai": "h1", "vi": "Visa EB-5", "en": "EB-5 Visa Guide"},
                                              {"id": "s002", "loai": "img", "vi": "Visa EB-5", "en": "Cover image for the EB-5 video"}])
            dich.lenh_nap([str(v)])
            tm = dich.doc_tm()
            self.assertEqual(len(tm), 2)  # không nuốt nhau vì trùng chữ Việt
            self.assertEqual(dich.goi_y_tm("Visa EB-5", tm, loai="h1"), "100%: EB-5 Visa Guide")
            self.assertEqual(dich.goi_y_tm("Visa EB-5", tm, loai="img"), "100%: Cover image for the EB-5 video")
        finally:
            dich.TM, dich.TMX = cu_tm, cu_tmx; shutil.rmtree(tmp)

    def test_bo_link_bo_the_giu_cho(self):
        self.assertEqual(dich.bo_link("{1}EB-5{/1} [x](y){br}{2/}"), "EB-5 [x] ")

    def test_tm_khop_chu_khac_the(self):
        tm = [{"vi": "Visa Định Cư Mỹ EB-5", "en": "EB-5 Visa"}]
        self.assertTrue(dich.goi_y_tm("Visa Định Cư Mỹ {1}EB-5{/1}", tm).startswith("100% (gắn lại thẻ {1} {/1}): "))
        self.assertEqual(dich.goi_y_tm("Visa định cư Mỹ EB-5", tm), "100%: EB-5 Visa")


MAU_JSON = GOC / "tests" / "mau" / "acf-product-mau.json"
DICH_JSON = {
    "Visa định cư Mỹ EB-5": "EB-5 Visa",
    "Visa Định Cư Mỹ {1}EB-5{/1}": "{1}EB-5{/1} Visa: A U.S. Green Card for Your Family",
    "EB-5 là gì?": "What is the EB-5 program?",
    "Mức đầu tư tối thiểu **800.000 USD**": "Minimum investment: **US$800,000**",
    "Yêu cầu tư vấn": "Book a consultation",
    "MỐC QUAN TRỌNG": "KEY DATE",
    "Hạn cuối nộp hồ sơ {1}I-526E{/1}:{br} {2}30 / 09 / 2026{/2}": "Form {1}I-526E{/1} filing deadline:{br} {2}09 / 30 / 2026{/2}",
    "Mốc **30/09/2026** ảnh hưởng tới hồ sơ. Xem [visa E-2](https://immgroup.com/dau-tu-dinh-cu-my/visa-my-e2/#faq).":
        "The **September 30, 2026** date affects your petition. See the [E-2 visa](https://immgroup.com/dau-tu-dinh-cu-my/visa-my-e2/#faq).",
    "Xét duyệt & quyền lợi": "Adjudication & benefits",
    "Gia đình tại Mỹ": 'A family in the "U.S."',
    "ĐIỀU KIỆN": "ELIGIBILITY",
    "Nguồn vốn hợp pháp": "Lawful source of funds",
    "Chứng minh nguồn gốc số tiền": "Documented source of funds",
    "TEA - khu vực việc làm mục tiêu.": "[BO]",
    "Tìm hiểu dịch vụ của IMM Group:": "Explore IMM Group's services:",
    "Visa Mỹ E-2": "U.S. E-2 Visa",
}


def khung(x):
    if isinstance(x, dict):
        return {k: khung(v) for k, v in x.items()}
    if isinstance(x, list):
        return [khung(v) for v in x]
    return type(x).__name__


class JsonAcf(unittest.TestCase):
    """Việc có nguồn là tệp xuất của ACF Page Importer: moi → dien → ghep ra tệp JSON import được."""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.cu = (dich.GOC, dich.TM, dich.LIEN_KET)
        dich.GOC, dich.TM, dich.LIEN_KET = self.tmp, self.tmp / "tm.jsonl", self.tmp / "lk.tsv"
        dich.LIEN_KET.write_text("vi\ten\tghi_chu\nhttps://immgroup.com/dau-tu-dinh-cu-my/visa-my-e2/\thttps://immgroup.com/en/e-2-visa/\t\n", encoding="utf-8")
        dich.lenh_moi([str(MAU_JSON)])
        self.viec = next((self.tmp / "viec").iterdir())

    def tearDown(self):
        dich.GOC, dich.TM, dich.LIEN_KET = self.cu
        shutil.rmtree(self.tmp)

    def dich_het(self, sua=None):
        dong = []
        for r in dich.doc_tsv(self.viec / "song-ngu.tsv"):
            en = (sua or {}).get(r["id"], DICH_JSON[r["vi"]])
            dong.append(f"[{r['id']}|#bo: chỉ giải nghĩa cho người Việt] [BO]" if en == "[BO]" else f"[{r['id']}] {en}")
        (self.tmp / "d.txt").write_text("\n".join(dong), encoding="utf-8")
        dich.lenh_dien([str(self.viec), str(self.tmp / "d.txt")])

    def test_tach_doan(self):
        r = dich.doc_tsv(self.viec / "song-ngu.tsv")
        self.assertEqual([(x["loai"], x["vi"]) for x in r], [
            ("tieu-de", "Visa định cư Mỹ EB-5"), ("h1", "Visa Định Cư Mỹ {1}EB-5{/1}"), ("h3", "EB-5 là gì?"),
            ("li", "Mức đầu tư tối thiểu **800.000 USD**"), ("cta", "Yêu cầu tư vấn"), ("p", "MỐC QUAN TRỌNG"),
            ("p", "Hạn cuối nộp hồ sơ {1}I-526E{/1}:{br} {2}30 / 09 / 2026{/2}"),
            ("p", "Mốc **30/09/2026** ảnh hưởng tới hồ sơ. Xem [visa E-2](https://immgroup.com/dau-tu-dinh-cu-my/visa-my-e2/#faq)."),
            ("h3", "Xét duyệt & quyền lợi"), ("img", "Gia đình tại Mỹ"), ("h2", "ĐIỀU KIỆN"), ("h3", "Nguồn vốn hợp pháp"),
            ("li", "Chứng minh nguồn gốc số tiền"), ("p", "TEA - khu vực việc làm mục tiêu."),
            ("p", "Tìm hiểu dịch vụ của IMM Group:"), ("li", "Visa Mỹ E-2")])
        self.assertEqual(r[6]["src"], "acf.product_important_content#2")
        self.assertTrue((self.viec / "nguon.json").exists())
        self.assertEqual(dich.doc_meta(self.viec)["bo_chuyen"], "acf-product-2026")

    def test_ghep_ra_json_import(self):
        self.dich_het()
        dich.lenh_ghep([str(self.viec)])
        goc = json.loads(MAU_JSON.read_text(encoding="utf-8"))[0]
        ra = json.loads((self.viec / "ban-giao" / "visa-mau.en.json").read_text(encoding="utf-8"))[0]
        self.assertEqual(khung(ra), khung(goc))  # cùng khoá, cùng số dòng repeater, cùng kiểu
        self.assertEqual((ra["post_id"], ra["post_slug"], ra["post_title"]), (123, "/visa-mau/", "EB-5 Visa"))
        a = ra["acf"]
        self.assertEqual(a["product_hero_title"]["product_hero_title_content"], '<span class=\\"accent\\">EB-5</span> Visa: A U.S. Green Card for Your Family')
        self.assertEqual(a["product_hero_content"][0]["product_hero_content_text"][0]["product_hero_content_text_item"], "Minimum investment: <strong>US$800,000</strong>")
        ic = a["product_important_content"]
        for mau in ['<!-- Khối mốc -->',
                    '<span class="material-symbols-outlined sm">alarm</span>\r\n    KEY DATE\r\n  </span>',
                    'Form <span style="color:var(--navy);">I-526E</span> filing deadline:<br> <span class="date-num">09 / 30 / 2026</span>\r\n  </div>',
                    '<strong style="color:var(--text-main);">September 30, 2026</strong>',
                    '<a href="https://immgroup.com/en/e-2-visa/#faq" target="_blank">E-2 visa</a>.',
                    '<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="10"/><path d="M9 12l2 2 4-4"/></svg>',
                    'Adjudication &amp; benefits', 'alt="A family in the &quot;U.S.&quot;"']:
            self.assertIn(mau, ic)
        self.assertEqual(a["product_conditions_items"][0]["product_conditions_item_footnote"], "")
        self.assertEqual(a["product_hero_cta"][0]["product_hero_cta_url"], "https://immgroup.com/dang-ky-tu-van/")
        self.assertEqual(a["product_care_service_links"][0]["product_care_service_link_url"], "https://immgroup.com/en/e-2-visa/")
        self.assertEqual(a["product_care_service_desc"], "<p>Explore IMM Group's services:</p>\n")
        for k in ("product_hero_video", "product_projects_posts", "product_faq_mode"):
            self.assertEqual(a[k], goc["acf"][k])
        self.assertEqual((a["product_hero_content"][0]["product_hero_content_icon"], a["product_hero_cta"][0]["product_hero_cta_is_primary"]), ("info", True))
        _, loi, canh = dich.dung_json(self.viec, dich.doc_tsv(self.viec / "song-ngu.tsv"))
        self.assertEqual(loi, [])
        self.assertTrue(any("dang-ky-tu-van" in c for c in canh) and any("97875" in c for c in canh), canh)
        md = (self.viec / "ban-giao" / "bai-dich.en.md").read_text(encoding="utf-8")
        self.assertIn("# EB-5 Visa: A U.S. Green Card for Your Family", md)
        self.assertNotIn("{1}", md)

    def test_the_lech_khong_ghi_json(self):
        self.dich_het({"s007": "Form {1}I-526E{/1} filing deadline:{br} {2}09 / 30 / 2026"})
        _, loi, _ = dich.dung_json(self.viec, dich.doc_tsv(self.viec / "song-ngu.tsv"))
        self.assertTrue(any("s007 THẺ HTML LỆCH" in x for x in loi), loi)
        with self.assertRaises(SystemExit):
            dich.lenh_ghep([str(self.viec)])
        self.assertFalse((self.viec / "ban-giao" / "visa-mau.en.json").exists())
        ma, bc = chay_kiem(self.viec)
        self.assertEqual(ma, 1)
        self.assertIn("THẺ HTML LỆCH", bc.split("## CẢNH BÁO")[0])

    def test_con_doan_chua_dich(self):
        self.dich_het({"s003": ""})
        _, loi, _ = dich.dung_json(self.viec, dich.doc_tsv(self.viec / "song-ngu.tsv"))
        self.assertTrue(any("chưa dịch" in x for x in loi), loi)

    def test_nguon_doi_sau_moi(self):
        p = self.viec / "nguon.json"
        p.write_text(p.read_text(encoding="utf-8").replace("Nguồn vốn hợp pháp", "Nguồn vốn sạch"), encoding="utf-8")
        _, loi, _ = dich.dung_json(self.viec, dich.doc_tsv(self.viec / "song-ngu.tsv"))
        self.assertIn("không còn khớp", loi[0])

    def test_vong_tron_giu_nguyen_html(self):
        goc = json.loads(MAU_JSON.read_text(encoding="utf-8"))[0]
        bc = dich.doc_bo_chuyen("acf-product-2026")
        ds = [d for d in dich.doan_json(goc, bc) if d["duong"] == ("acf", "product_important_content")]
        chuoi = goc["acf"]["product_important_content"]
        ra = dich.ghep_truong_html(chuoi, [(d, d["vi"]) for d in ds])
        self.assertEqual(dich.gon(ra), dich.gon(chuoi))  # chỉ khác xuống dòng BÊN TRONG đoạn (gộp thành 1 dấu cách)
        self.assertIn("alarm</span>\r\n    MỐC QUAN TRỌNG\r\n  </span>", ra)  # khoảng trắng ở mép đoạn giữ nguyên

    def test_json_dan_vao_chat_luu_nham_duoi_md(self):
        p = self.tmp / "dan-tu-chat.md"
        p.write_text(MAU_JSON.read_text(encoding="utf-8"), encoding="utf-8")
        dich.lenh_moi([str(p), "--ten", "dan-chat"])
        viec = next((self.tmp / "viec").glob("*-dan-chat"))
        self.assertEqual(dich.doc_meta(viec)["dang"], "acf-json")
        self.assertEqual(dich.doc_tsv(viec / "song-ngu.tsv")[1]["vi"], "Visa Định Cư Mỹ {1}EB-5{/1}")

    def test_moi_trung_ten_khong_ghi_de(self):
        self.dich_het()
        dich.lenh_moi([str(MAU_JSON)])
        cu, moi = sorted((self.tmp / "viec").iterdir())
        self.assertEqual((cu.name[11:], moi.name[11:]), ("visa-mau", "visa-mau-2"))
        self.assertTrue(all(r["en"] for r in dich.doc_tsv(cu / "song-ngu.tsv")))  # bản dịch việc cũ còn nguyên

    def test_nhieu_trang_phai_chon(self):
        nhieu = self.tmp / "nhieu.json"
        mot = json.loads(MAU_JSON.read_text(encoding="utf-8"))[0]
        nhieu.write_text(json.dumps([mot, dict(mot, post_slug="/khac/", post_id=9)], ensure_ascii=False), encoding="utf-8")
        with self.assertRaises(SystemExit):
            dich.lenh_moi([str(nhieu)])
        dich.lenh_moi([str(nhieu), "--trang", "/khac/", "--ten", "khac"])
        self.assertEqual(dich.doc_meta(next((self.tmp / "viec").glob("*-khac")))["post_id"], 9)


class GanLink(unittest.TestCase):
    """ganlink: gắn link bản dịch trên GitHub vào bảng đối chiếu xlsx."""

    def setUp(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("máy chưa có openpyxl")
        self.tmp = Path(tempfile.mkdtemp())
        self.viec = self.tmp / "viec" / "2026-09-22-trang-thu"
        (self.viec / "ban-giao").mkdir(parents=True)
        (self.viec / "meta.json").write_text(json.dumps({"post_slug": "/trang-thu/"}), encoding="utf-8")
        for t in ("en", "vi"):
            (self.viec / "ban-giao" / f"bai-dich.{t}.md").write_text("# x\n", encoding="utf-8")
        self.xlsx = self.tmp / "bang.xlsx"
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = "translate new"
        ws.append(["link", "vi", "en"])
        ws.append(["https://immgroup.com/muc-cha/trang-thu/", "x", None])
        ws.append(["https://immgroup.com/muc-khac/trang-khac/", None, None])
        wb.save(self.xlsx)
        self.bang_cu = dich.BANG_BAN_DICH
        dich.BANG_BAN_DICH = self.tmp / "ban-dich-vi-en.tsv"

    def tearDown(self):
        dich.BANG_BAN_DICH = self.bang_cu
        shutil.rmtree(self.tmp)

    def o(self, dong, cot):
        import openpyxl
        return openpyxl.load_workbook(self.xlsx)["translate new"].cell(dong, cot).value

    def test_gan_dung_dong_theo_slug(self):
        dich.lenh_ganlink([str(self.viec), "--xlsx", str(self.xlsx), "--repo", "https://ví-dụ/blob/main"])
        self.assertEqual(self.o(2, 2), "https://ví-dụ/blob/main/viec/2026-09-22-trang-thu/ban-giao/bai-dich.vi.md")
        self.assertEqual(self.o(2, 3), "https://ví-dụ/blob/main/viec/2026-09-22-trang-thu/ban-giao/bai-dich.en.md")
        self.assertIsNone(self.o(3, 2))  # dòng khác không bị đụng
        self.assertIn("2026-09-22-trang-thu", dich.BANG_BAN_DICH.read_text(encoding="utf-8"))

    def test_thu_khong_ghi(self):
        dich.lenh_ganlink([str(self.viec), "--xlsx", str(self.xlsx), "--thu"])
        self.assertEqual(self.o(2, 2), "x")  # giữ nguyên ô cũ
        self.assertFalse(dich.BANG_BAN_DICH.exists())

    def test_khong_khop_slug_thi_bao_loi_va_khong_ghi(self):
        (self.viec / "meta.json").write_text(json.dumps({"post_slug": "/khong-co-trong-bang/"}), encoding="utf-8")
        with self.assertRaises(SystemExit):
            dich.lenh_ganlink([str(self.viec), "--xlsx", str(self.xlsx)])
        self.assertEqual(self.o(2, 2), "x")

    def test_chua_ghep_thi_khong_gan(self):
        (self.viec / "ban-giao" / "bai-dich.vi.md").unlink()
        with self.assertRaises(SystemExit):
            dich.lenh_ganlink([str(self.viec), "--xlsx", str(self.xlsx)])
        self.assertEqual(self.o(2, 2), "x")

    def test_chi_ghi_tsv_khi_khong_neu_xlsx(self):
        dich.lenh_ganlink([str(self.viec), "--link", "https://immgroup.com/muc-cha/trang-thu/"])
        d = list(csv.DictReader(dich.BANG_BAN_DICH.open(encoding="utf-8"), delimiter="\t"))
        self.assertEqual(d[0]["link"], "https://immgroup.com/muc-cha/trang-thu/")
        self.assertTrue(d[0]["en"].endswith("bai-dich.en.md"))
        self.assertEqual(self.o(2, 2), "x")  # tệp Excel của team không bị đụng

    def test_lan_sau_tu_nho_url_khong_can_link(self):
        dich.lenh_ganlink([str(self.viec), "--link", "https://immgroup.com/muc-cha/trang-thu/"])
        dich.lenh_ganlink([str(self.viec)])  # không nêu --link, không nêu --xlsx
        d = list(csv.DictReader(dich.BANG_BAN_DICH.open(encoding="utf-8"), delimiter="\t"))
        self.assertEqual(len(d), 1)

    def test_chua_biet_url_thi_bao_loi(self):
        with self.assertRaises(SystemExit):
            dich.lenh_ganlink([str(self.viec)])  # bảng trống, không có --link, không có --xlsx
        self.assertFalse(dich.BANG_BAN_DICH.exists())

    def test_xuat_bang_excel_moi(self):
        import openpyxl
        dich.lenh_ganlink([str(self.viec), "--link", "https://immgroup.com/muc-cha/trang-thu/"])
        ra = self.tmp / "gui-team.xlsx"
        dich.lenh_ganlink(["--xuat", str(ra)])
        ws = openpyxl.load_workbook(ra).active
        self.assertEqual([c.value for c in ws[1]], ["link", "vi", "en", "viec", "ngay"])
        self.assertEqual(ws.cell(2, 1).value, "https://immgroup.com/muc-cha/trang-thu/")
        self.assertEqual(ws.cell(2, 4).value, "2026-09-22-trang-thu")

    def test_ep_link_khi_slug_lech(self):
        (self.viec / "meta.json").write_text(json.dumps({"post_slug": "/slug-lech/"}), encoding="utf-8")
        dich.lenh_ganlink([str(self.viec), "--xlsx", str(self.xlsx),
                           "--link", "https://immgroup.com/muc-cha/trang-thu/"])
        self.assertTrue(str(self.o(2, 3)).endswith("bai-dich.en.md"))


if __name__ == "__main__":
    unittest.main()

#!/usr/bin/env python3
"""Bộ công cụ dịch web IMM Group (Việt -> Anh). Chỉ dùng thư viện chuẩn + bs4 + python-docx (+ openpyxl cho lệnh ganlink).

Lệnh:
  moi    <url|tệp> [--ten slug]   tạo thư mục việc, tách đoạn ra song-ngu.tsv, gợi ý từ bộ nhớ dịch
         tệp .json xuất từ ACF Page Importer: [--trang <số|post_slug|post_id>] [--bo-chuyen <tên>]
  tukhoa <thư-mục-việc> "hạt 1" "hạt 2" [--thi-truong us,gb,au,ca,sg,ae]
  tygia  <thư-mục-việc>           lấy tỷ giá USD/VND, USD/EUR... ghi ty-gia.json
  ghep   <thư-mục-việc> [--chi-vi]  dựng ban-giao/bai-dich.en.md + .html + bai-dich.vi.md từ song-ngu.tsv
         (việc JSON: thêm ban-giao/<slug>.en.json · --chi-vi: chỉ dựng lại bản VI, không đụng bản EN đã qua cửa 0)
  kiem   <thư-mục-việc>           CỬA 0 bằng máy — thoát mã 1 nếu có LỖI
  nap    <thư-mục-việc>           nạp câu đã chốt vào bộ nhớ dịch (trùng thì xoá cũ giữ mới) + xuất TMX
  kiemtn                          kiểm bảng thuật ngữ (trùng, thiếu cột)
  xem    <thư-mục-việc> [--tu s001] [--den s150] [--anh]   in gọn nguồn (và bản Anh nếu --anh) để đọc/dịch
  dien   <thư-mục-việc> <tệp.txt>... [--tm100] [--trung]  điền bản dịch/bản sửa dạng "[s001] text" hoặc "[s008|#bo: lý do] [BO]"
  ganlink <thư-mục-việc>... | --tat-ca   ghi link bản dịch trên GitHub vào lien-ket/ban-dich-vi-en.tsv
         [--link <url trang gốc>] ép dòng cho MỘT việc khi slug không khớp (lần sau tự nhớ) · --thu = chỉ xem
         [--xlsx <tệp.xlsx>] ghi thêm vào MỘT bảng Excel có sẵn (trang "translate new", cột link/vi/en) — cần openpyxl
  ganlink --xuat <tệp.xlsx>       dựng bảng Excel MỚI từ ban-dich-vi-en.tsv để gửi team (không đụng tệp của team)
"""
import copy, csv, datetime as dt, difflib, fnmatch, html, json, re, subprocess, sys, time, unicodedata
import urllib.parse, urllib.request
from pathlib import Path

GOC = Path(__file__).resolve().parent.parent
TM = GOC / "bo-nho-dich" / "bo-nho-dich.jsonl"
TMX = GOC / "bo-nho-dich" / "bo-nho-dich.tmx"
THUAT_NGU = GOC / "thuat-ngu" / "thuat-ngu.csv"
BO_CHUYEN = GOC / "bo-chuyen"
LIEN_KET = GOC / "lien-ket" / "lien-ket-vi-en.tsv"
BANG_BAN_DICH = GOC / "lien-ket" / "ban-dich-vi-en.tsv"
REPO_BLOB = "https://github.com/kelvinimm/web-imm-translation/blob/main"
SHEET_DOI_CHIEU = "translate new"
COT =["id", "loai", "vi", "en", "tm", "ghi_chu", "src"]
UA = "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 Chrome/126 Safari/537.36"
KHOI = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote", "figcaption", "td", "th", "dt", "dd", "button"]
BO_QUA_TAG = ["script", "style", "noscript", "svg", "form", "nav", "header", "footer", "iframe", "select", "option", "template"]
BO_QUA_CLASS = re.compile(r"footer|breadcrumb|share|social|popup|modal|cookie|site-header|menu|material-(icons|symbols)|dashicons|\\bfa-|screen-reader", re.I)
TEN_ICON = re.compile(r"^[a-z]+(?:_[a-z]+)*$")
VN_CHU = re.compile(r"[ăâđêôơưàáảãạằắẳẵặầấẩẫậèéẻẽẹềếểễệìíỉĩịòóỏõọồốổỗộờớởỡợùúủũụừứửữựỳýỷỹỵ]", re.I)
# chữ CHỈ tiếng Việt mới có (é, í, ã… có cả trong tên Bồ Đào Nha, Tây Ban Nha, Pháp)
VN_RIENG = re.compile(r"[ăđơưạảấầẩẫậắằẳẵặẹẻẽếềểễệỉịọỏốồổỗộớờởỡợụủứừửữựỳỵỷỹ]", re.I)


def con_tieng_viet(s):
    """Có chữ riêng tiếng Việt, hoặc ≥2 từ mang dấu (vd "khách hàng") thì coi là còn tiếng Việt."""
    if VN_RIENG.search(s):
        return True
    # từ tiếng Việt là âm tiết ngắn; tên Tây có dấu (André, Jurídico) thường đứng lẻ hoặc dài hơn
    return sum(1 for w in re.findall(r"[\wÀ-ỹ]+", s) if VN_CHU.search(w) and len(w) <= 5) >= 2


# ---------- tiện ích ----------
def lay_url(url, timeout=25):
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept-Language": "vi,en;q=0.8"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read().decode(r.headers.get_content_charset() or "utf-8", "replace")


def slugify(s, toi_da=60):
    s = unicodedata.normalize("NFD", s.replace("đ", "d").replace("Đ", "D"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn").lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:toi_da].strip("-") or "bai"


def chuan(s):
    return re.sub(r"\s+", " ", (s or "")).strip().lower()


def doc_tsv(p):
    with open(p, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def ghi_tsv(p, dong):
    with open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, COT, delimiter="\t", extrasaction="ignore")
        w.writeheader()
        for d in dong:
            w.writerow({k: re.sub(r"[\t\r\n]+", " ", d.get(k, "") or "") for k in COT})


def doc_tm():
    if not TM.exists():
        return []
    return [json.loads(l) for l in TM.read_text(encoding="utf-8").splitlines() if l.strip()]


def doc_thuat_ngu():
    if not THUAT_NGU.exists():
        return []
    with open(THUAT_NGU, encoding="utf-8", newline="") as f:
        return [r for r in csv.DictReader(f) if (r.get("vi") or r.get("en"))]


# ---------- tách đoạn ----------
def inline_md(el):
    from bs4 import NavigableString, Comment
    out = []
    for c in el.children:
        if isinstance(c, Comment):
            continue
        if isinstance(c, NavigableString):
            out.append(str(c))
        elif c.name == "br":
            out.append(" ")
        elif c.name == "a" and c.get("href") and c.get_text(strip=True):
            out.append(f"[{c.get_text(' ', strip=True)}]({c['href']})")
        elif c.name in ("strong", "b") and c.get_text(strip=True):
            out.append(f"**{c.get_text(' ', strip=True)}**")
        elif c.name in BO_QUA_TAG:
            continue
        else:
            out.append(inline_md(c))
    return re.sub(r"\s+", " ", "".join(out)).strip()


def tach_html(src_html):
    from bs4 import BeautifulSoup, NavigableString, Comment
    s = BeautifulSoup(src_html, "html.parser")
    for cm in s.find_all(string=lambda t: isinstance(t, Comment)):
        cm.extract()
    meta = {"title": s.title.get_text(strip=True) if s.title else ""}
    for k, sel in [("description", {"name": "description"}), ("og_title", {"property": "og:title"})]:
        m = s.find("meta", attrs=sel)
        meta[k] = m.get("content", "") if m else ""
    for l in s.find_all("link", rel="alternate"):
        if l.get("hreflang"):
            meta["hreflang_" + l["hreflang"]] = l.get("href", "")
    c = s.find("link", rel="canonical")
    meta["canonical"] = c.get("href", "") if c else ""
    goc = s.select_one("#content") or s.find("main") or s.find("article") or s.body or s
    for t in goc.find_all(BO_QUA_TAG):
        t.decompose()
    for t in goc.find_all(True):
        if t.attrs is not None and BO_QUA_CLASS.search(" ".join(t.get("class", []) or []) + " " + (t.get("id") or "")):
            t.decompose()
    dong = []

    def them(loai, vi, src=""):
        vi = re.sub(r"\s+", " ", vi or "").strip()
        vi = re.sub(r"\s+(arrow_forward|arrow_back|chevron_right|expand_more|east|north_east)$", "", vi)
        if loai != "img" and TEN_ICON.fullmatch(vi) and ("_" in vi or vi in ("info", "school", "alarm", "check", "close", "menu", "search", "home", "star")):
            return
        if vi or loai == "img":
            dong.append({"loai": loai, "vi": vi, "src": src})

    def co_khoi_con(el):
        return el.find(KHOI + ["img"]) is not None

    def duyet(el):
        tam = []
        for c in el.children:
            if isinstance(c, NavigableString):
                if c.strip():
                    tam.append(str(c))
                continue
            if c.name is None:
                continue
            if c.name == "img":
                them("p", "".join(tam)); tam = []
                them("img", c.get("alt", ""), c.get("data-src") or c.get("src", ""))
            elif c.name in KHOI and not co_khoi_con(c):
                them("p", "".join(tam)); tam = []
                them("quote" if c.name == "blockquote" else ("cta" if c.name == "button" else c.name), inline_md(c))
            elif c.name in ("a", "strong", "b", "span", "em", "i") and not co_khoi_con(c):
                tam.append(inline_md(c) if c.name != "span" else c.get_text(" "))
            else:
                them("p", "".join(tam)); tam = []
                duyet(c)
        them("p", "".join(tam))

    duyet(goc)
    return meta, dong


def tach_md(text):
    dong, doan = [], []

    def xa():
        if doan:
            dong.append({"loai": "p", "vi": " ".join(doan).strip(), "src": ""}); doan.clear()

    for l in text.splitlines():
        l = l.rstrip()
        m_h = re.match(r"^(#{1,6})\s+(.*)", l)
        m_img = re.match(r"^!\[(.*?)\]\((.*?)\)\s*$", l)
        m_li = re.match(r"^\s*(?:[-*+]|\d+[.)])\s+(.*)", l)
        if not l.strip():
            xa()
        elif m_h:
            xa(); dong.append({"loai": f"h{len(m_h.group(1))}", "vi": m_h.group(2).strip(), "src": ""})
        elif m_img:
            xa(); dong.append({"loai": "img", "vi": m_img.group(1), "src": m_img.group(2)})
        elif m_li:
            xa(); dong.append({"loai": "li", "vi": m_li.group(1).strip(), "src": ""})
        elif l.startswith(">"):
            xa(); dong.append({"loai": "quote", "vi": l.lstrip("> ").strip(), "src": ""})
        else:
            doan.append(l.strip())
    xa()
    return dong


def tach_docx(p):
    import docx
    dong = []
    for para in docx.Document(p).paragraphs:
        t = para.text.strip()
        if not t:
            continue
        st = (para.style.name or "").lower()
        m = re.search(r"heading\s*(\d)", st)
        if st == "title":
            dong.append({"loai": "h1", "vi": t, "src": ""})
        elif m:
            dong.append({"loai": f"h{m.group(1)}", "vi": t, "src": ""})
        elif "list" in st:
            dong.append({"loai": "li", "vi": t, "src": ""})
        else:
            dong.append({"loai": "p", "vi": t, "src": ""})
    return dong


def khoa_tm(e):
    # alt ảnh và chữ thường cùng câu Việt vẫn dịch khác nhau → hai mục riêng trong bộ nhớ dịch
    return (e.get("loai") == "img", chuan(e["vi"]))


def goi_y_tm(vi, tm, nguong=0.85, loai=""):
    # so khớp bỏ qua thẻ giữ chỗ {1}…{/1}; khớp chữ mà khác thẻ thì không cho `dien --tm100` tự lấp
    k = chuan(bo_the(vi))
    if not k or len(k) < 3:
        return ""
    tot = None
    for e in tm:
        if loai and (e.get("loai") == "img") != (loai == "img"):
            continue
        kk = chuan(bo_the(e["vi"]))
        if kk == k:
            if THE_RE.findall(vi) != THE_RE.findall(e["vi"]):
                return f"100% (gắn lại thẻ {' '.join(THE_RE.findall(vi))}): {e['en']}"
            return f"100%: {e['en']}"
        sm = difflib.SequenceMatcher(None, k, kk)
        if sm.real_quick_ratio() < nguong or sm.quick_ratio() < nguong:
            continue
        r = sm.ratio()
        if r >= nguong and (not tot or r > tot[0]):
            tot = (r, e["en"])
    return f"{int(tot[0] * 100)}%: {tot[1]}" if tot else ""


# ---------- JSON xuất từ ACF Page Importer ----------
# Đoạn dịch lấy từ các trường chữ theo bo-chuyen/<tên>.json. Trường có HTML được chia theo khối; thẻ trong câu
# hiện dạng **đậm**, [chữ](link), {1}…{/1} (thẻ khác), {2/} (thẻ lẻ), {br} — bản dịch phải giữ đủ thẻ.
THE_RE = re.compile(r"\{/?\d+/?\}|\{br\}")
TOKEN_RE = re.compile(r"<!--.*?-->|<![^>]*>|</?[a-zA-Z][^>]*>", re.S)
THE_KHOI = {"address", "article", "aside", "blockquote", "button", "dd", "details", "dialog", "div", "dl", "dt",
            "fieldset", "figcaption", "figure", "footer", "form", "h1", "h2", "h3", "h4", "h5", "h6", "header", "hr",
            "li", "main", "nav", "ol", "p", "pre", "section", "summary", "table", "tbody", "td", "tfoot", "th",
            "thead", "tr", "ul"}
THE_LE = {"br", "img", "wbr", "input", "hr", "source", "col", "area", "embed", "track"}
THE_KIN = {"svg", "script", "style", "noscript", "iframe", "template", "select", "textarea", "video", "audio",
           "canvas", "object", "math"}  # bỏ nguyên khối, không dịch
LOP_ICON = re.compile(r"material-(?:icons|symbols)|dashicons|\bfa-", re.I)
LOAI_KHOI = {"li": "li", "dt": "li", "dd": "li", "button": "cta", "blockquote": "quote",
             **{f"h{i}": f"h{i}" for i in range(1, 7)}}
MIEN_IMM = ("immgroup.com", "www.immgroup.com")
LAP_RE = re.compile(r"\*\*|\[((?:[^\[\]]|\[[^\]]*\])*)\]\(([^)\s]*)\)|\{(\d+)\}|\{/(\d+)\}|\{(\d+)/\}|\{br\}")


def bo_the(s):
    return THE_RE.sub(lambda m: " " if m.group(0) == "{br}" else "", s or "")


def the_giu_cho(s):
    return sorted(re.findall(r"\{/?\d+/?\}", s or ""))


def gon(s):
    return re.sub(r"\s+", " ", s or "").strip()


def la_html(s):
    return bool(re.search(r"</?[a-zA-Z][^>]*>|&(?:[a-zA-Z]+|#\d+|#x[0-9a-fA-F]+);", s))


def tach_token(s):
    """Chuỗi HTML → token (kiểu, chuỗi gốc, tên thẻ); kiểu: chu · cm · mo · tat · le. Ghép lại các token = chuỗi gốc."""
    ra, vt = [], 0
    for m in TOKEN_RE.finditer(s):
        if m.start() > vt:
            ra.append(("chu", s[vt:m.start()], ""))
        t = m.group(0)
        if t.startswith("<!"):
            ra.append(("cm", t, ""))
        else:
            ten = re.match(r"</?([a-zA-Z][\w:-]*)", t).group(1).lower()
            ra.append(("tat" if t.startswith("</") else ("le" if t.rstrip().endswith("/>") or ten in THE_LE else "mo"), t, ten))
        vt = m.end()
    if vt < len(s):
        ra.append(("chu", s[vt:], ""))
    return ra


def thuoc_tinh(raw, ten):
    # giá trị thuộc tính, chấp cả nháy đã escape (class=\"accent\") từ dữ liệu cũ
    return re.search(r"(?<![\w-])" + ten + r"\s*=\s*(\\?[\"'])(.*?)\1", raw, re.S)


def cap_the(tokens, tu, den):
    cap, ngan = {}, []
    for i in range(tu, den + 1):
        kieu, _, ten = tokens[i]
        if kieu == "mo":
            ngan.append(i)
        elif kieu == "tat":
            for j in range(len(ngan) - 1, -1, -1):
                if tokens[ngan[j]][2] == ten:
                    cap[ngan[j]], cap[i] = i, ngan[j]
                    del ngan[j:]
                    break
    return cap


def het_the(tokens, i):
    ten, sau = tokens[i][2], 0
    for j in range(i, len(tokens)):
        if tokens[j][2] == ten and tokens[j][0] == "mo":
            sau += 1
        elif tokens[j][2] == ten and tokens[j][0] == "tat":
            sau -= 1
            if sau == 0:
                return j
    return len(tokens) - 1


def chia_html(tokens, loai_goc):
    """Chia token HTML thành đoạn dịch: {"cach": "chay", "tu", "den", "loai"} hoặc {"cach": "alt", "i", "loai": "img"}."""
    kq, chay, khoi = [], [], []

    def xa():
        chu = [i for i in chay if tokens[i][0] == "chu" and html.unescape(tokens[i][1]).strip()]
        if chu and any(re.search(r"\w", html.unescape(tokens[i][1])) for i in chu):
            tu, den = chu[0], chu[-1]
            cap, doi = cap_the(tokens, chay[0], chay[-1]), True
            while doi:  # kéo vào đoạn thẻ nào có một đầu nằm trong đoạn
                doi = False
                for i in range(tu, den + 1):
                    j = cap.get(i)
                    if j is not None and not tu <= j <= den:
                        tu, den, doi = min(tu, j), max(den, j), True
            kq.append({"cach": "chay", "tu": tu, "den": den, "loai": LOAI_KHOI.get(khoi[-1], "p") if khoi else loai_goc})
        chay.clear()

    i = 0
    while i < len(tokens):
        kieu, raw, ten = tokens[i]
        lop = thuoc_tinh(raw, "class") if kieu == "mo" else None
        if kieu == "cm":
            xa()
        elif kieu == "mo" and (ten in THE_KIN or (lop and LOP_ICON.search(lop.group(2)))):
            xa(); i = het_the(tokens, i)
        elif ten in THE_KHOI:
            xa()
            if kieu == "mo":
                khoi.append(ten)
            elif kieu == "tat" and ten in khoi:
                del khoi[len(khoi) - 1 - khoi[::-1].index(ten):]
        elif ten == "img":
            xa()
            m = thuoc_tinh(raw, "alt")
            if m and re.search(r"\w", html.unescape(m.group(2))):
                kq.append({"cach": "alt", "i": i, "loai": "img"})
        else:
            chay.append(i)
        i += 1
    xa()
    return kq


def sang_doan(tokens, tu, den):
    """Token [tu, den] → chữ đoạn dịch + bảng thẻ gốc để dựng lại (tt)."""
    cap = cap_the(tokens, tu, den)
    tt, ra, n, ban = {"b": [], "a": [], "br": [], "ph": {}}, [], 0, {}
    for i in range(tu, den + 1):
        kieu, raw, ten = tokens[i]
        mo = raw if kieu == "mo" else (tokens[cap[i]][1] if i in cap else "")
        if kieu == "chu":
            ra.append(html.unescape(raw))
        elif ten == "br":
            tt["br"].append(raw); ra.append("{br}")
        elif ten in ("strong", "b") and i in cap:
            if kieu == "mo":
                tt["b"].append((raw, tokens[cap[i]][1]))
            ra.append("**")
        elif ten == "a" and i in cap and thuoc_tinh(mo, "href"):
            if kieu == "mo":
                tt["a"].append((raw, tokens[cap[i]][1])); ban[cap[i]] = html.unescape(thuoc_tinh(raw, "href").group(2)); ra.append("[")
            else:
                ra.append(f"]({ban[i]})")
        elif kieu == "mo" and i in cap:
            n += 1; tt["ph"][n] = (raw, tokens[cap[i]][1]); ban[cap[i]] = n; ra.append(f"{{{n}}}")
        elif kieu == "tat" and i in cap:
            ra.append(f"{{/{ban[i]}}}")
        else:
            n += 1; tt["ph"][n] = (raw, None); ra.append(f"{{{n}/}}")
    return gon("".join(ra)), tt


def sang_html(en, tt, doi_url=lambda u: u):
    """Bản dịch (có **, [](), {n}, {br}) → HTML, dùng lại đúng thẻ gốc; href đi qua doi_url."""
    trang = {"b": 0, "a": 0, "br": 0, "mo_b": []}

    def dung(s):
        ra, vt = [], 0
        for m in LAP_RE.finditer(s):
            ra.append(html.escape(s[vt:m.start()], quote=False)); vt = m.end()
            g = m.group(0)
            if g == "**":
                if trang["mo_b"]:
                    ra.append(trang["mo_b"].pop())
                else:
                    mo, tat = tt["b"][trang["b"]] if trang["b"] < len(tt["b"]) else ("<strong>", "</strong>")
                    trang["b"] += 1; ra.append(mo); trang["mo_b"].append(tat)
            elif g == "{br}":
                ra.append(tt["br"][min(trang["br"], len(tt["br"]) - 1)] if tt["br"] else "<br>"); trang["br"] += 1
            elif m.group(2) is not None:
                mo, tat = tt["a"][trang["a"]] if trang["a"] < len(tt["a"]) else ('<a href="">', "</a>")
                trang["a"] += 1
                h = thuoc_tinh(mo, "href")
                if h:
                    mo = mo[:h.start(2)] + html.escape(doi_url(m.group(2)), quote=False) + mo[h.end(2):]
                ra.append(mo + dung(m.group(1)) + tat)
            else:
                so = int(m.group(3) or m.group(4) or m.group(5))
                mo, tat = tt["ph"][so]
                ra.append(tat if m.group(4) else mo)
        ra.append(html.escape(s[vt:], quote=False))
        return "".join(ra)

    return dung(en) + "".join(reversed(trang["mo_b"]))


def duyet_la(obj, duong, ten=""):
    """Sinh (đường dẫn, tên trường gần nhất, giá trị) cho mọi lá; mảng số (ID bài) là một lá."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            yield from duyet_la(v, duong + (k,), k)
    elif isinstance(obj, list) and obj and not all(isinstance(x, int) and not isinstance(x, bool) for x in obj):
        for i, v in enumerate(obj):
            yield from duyet_la(v, duong + (i,), ten)
    else:
        yield duong, ten, obj


def chuoi_duong(duong):
    s = ""
    for k in duong:
        s += f"[{k}]" if isinstance(k, int) else ("." if s else "") + k
    return s


def lay_theo_duong(obj, duong):
    for k in duong:
        obj = obj[k]
    return obj


def dat_theo_duong(obj, duong, v):
    lay_theo_duong(obj, duong[:-1])[duong[-1]] = v


def luat_truong(ten, bc):
    for mau, kieu in bc.get("quy_tac", []):
        if fnmatch.fnmatchcase(ten, mau):
            return kieu
    return bc.get("mac_dinh", "p")


def doc_bo_chuyen(ten):
    p = BO_CHUYEN / f"{ten}.json"
    if not p.exists():
        sys.exit(f"Không có bộ chuyển {p}")
    bc = json.loads(p.read_text(encoding="utf-8"))
    bc["ten"] = ten
    return bc


def chon_bo_chuyen(acf, ten=""):
    if ten:
        return doc_bo_chuyen(ten)
    khop = [doc_bo_chuyen(p.stem) for p in sorted(BO_CHUYEN.glob("*.json"))]
    khop = [bc for bc in khop if set(acf) <= set(bc.get("truong_goc", []))]
    if len(khop) == 1:
        return khop[0]
    if not khop:
        sys.exit("Chưa có bộ chuyển cho template này. Trường gốc trong tệp: " + ", ".join(acf)
                 + f"\n→ tạo {BO_CHUYEN}/<tên>.json theo mẫu acf-product-2026.json rồi chạy lại")
    sys.exit("Nhiều bộ chuyển cùng khớp: " + ", ".join(bc["ten"] for bc in khop) + " — chọn bằng --bo-chuyen <tên>")


def doan_json(trang, bc):
    """Các đoạn cần dịch của một trang theo thứ tự trong tệp, kèm thông tin để dựng lại."""
    ds = []
    if isinstance(trang.get("post_title"), str) and re.search(r"\w", trang["post_title"]):
        ds.append({"loai": "tieu-de", "vi": gon(html.unescape(trang["post_title"])), "src": "post_title",
                   "duong": ("post_title",), "cach": "chu"})
    for duong, ten, v in duyet_la(trang.get("acf") or {}, ("acf",)):
        kieu = luat_truong(ten, bc)
        if kieu in ("bo", "url", "id") or not isinstance(v, str) or not re.search(r"\w", v):
            continue
        if not la_html(v):
            ds.append({"loai": kieu, "vi": gon(v), "src": chuoi_duong(duong), "duong": duong, "cach": "chu"})
            continue
        tokens = tach_token(v)
        for so, p in enumerate(chia_html(tokens, kieu), 1):
            d = dict(p, duong=duong, src=f"{chuoi_duong(duong)}#{so}")
            if p["cach"] == "alt":
                d.update(vi=gon(html.unescape(thuoc_tinh(tokens[p["i"]][1], "alt").group(2))), src=d["src"] + "@alt")
            else:
                d["vi"], d["tt"] = sang_doan(tokens, p["tu"], p["den"])
            ds.append(d)
    return ds


def la_json_acf(p):
    """Tệp (đuôi bất kỳ) chứa JSON xuất của ACF Page Importer."""
    try:
        s = p.read_text(encoding="utf-8").lstrip()
        return s[:1] in "[{" and '"acf"' in s and bool(json.loads(s))
    except (ValueError, OSError):
        return False


def tach_json_tep(p, args):
    """Đọc tệp xuất của ACF Page Importer → (đoạn, meta, trang đã chọn, tên việc)."""
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except ValueError as e:
        sys.exit(f"Tệp JSON lỗi cú pháp: {e}")
    data = [data] if isinstance(data, dict) else data
    if not (isinstance(data, list) and data and all(isinstance(x, dict) and isinstance(x.get("acf"), dict) for x in data)):
        sys.exit('Tệp JSON không đúng dạng xuất của ACF Page Importer (mảng các trang, mỗi trang có "acf")')
    chon = args[args.index("--trang") + 1] if "--trang" in args else ("0" if len(data) == 1 else "")
    tim = [x for i, x in enumerate(data) if chon in (str(i), x.get("post_slug"), str(x.get("post_id", "")))]
    if len(tim) != 1:
        sys.exit("Tệp có nhiều trang — mỗi việc dịch 1 trang, chọn bằng --trang <số|post_slug|post_id>:\n"
                 + "\n".join(f"  {i}: {x.get('post_slug')} · {x.get('post_title')}" for i, x in enumerate(data)))
    trang = tim[0]
    bc = chon_bo_chuyen(trang["acf"], args[args.index("--bo-chuyen") + 1] if "--bo-chuyen" in args else "")
    dong = [{k: d[k] for k in ("loai", "vi", "src")} for d in doan_json(trang, bc)]
    meta = {"dang": "acf-json", "bo_chuyen": bc["ten"], "post_slug": trang.get("post_slug", ""),
            "post_title": trang.get("post_title", "")}
    if "post_id" in trang:
        meta["post_id"] = trang["post_id"]
    else:
        print("⚠ Tệp không có post_id (export cũ): Tools → Dịch trang ACF sẽ TỪ CHỐI nhập tệp dịch ra từ đây. "
              "Export lại TRANG TIẾNG ANH bằng Tools → Dịch trang ACF rồi chạy `moi` tệp đó (bản dịch đã làm lấy lại qua bộ nhớ dịch).")
    ten = slugify((trang.get("post_slug") or "").strip("/").split("/")[-1] or trang.get("post_title", ""))
    return dong, meta, trang, ten


def khoa_lien_ket(u):
    """Khoá tra bảng liên kết: 'id:123' cho ID bài, đường dẫn '/a/b/' cho trang immgroup.com, None cho thứ khác."""
    u = (u or "").strip()
    if re.fullmatch(r"\d+", u):
        return "id:" + u
    p = urllib.parse.urlparse(u)
    duong = p.path or ("/" if p.netloc else "")
    if p.scheme not in ("", "http", "https") or (p.netloc and p.netloc.lower() not in MIEN_IMM) or not duong.startswith("/"):
        return None
    return duong.rstrip("/").lower() + "/"


def doc_lien_ket():
    if not LIEN_KET.exists():
        return {}
    return {khoa_lien_ket(r["vi"]): r["en"].strip() for r in doc_tsv(LIEN_KET)
            if khoa_lien_ket(r.get("vi")) and (r.get("en") or "").strip()}


def doi_lien_ket(u, bang):
    """Đổi link trang Việt sang trang Anh theo bảng. Trả (link mới, còn trỏ trang Việt?)."""
    k = khoa_lien_ket(u)
    if k in bang:
        p = urllib.parse.urlparse(u)
        return bang[k] + (f"?{p.query}" if p.query else "") + (f"#{p.fragment}" if p.fragment else ""), False
    return u, bool(k) and not k.startswith(("id:", "/en/", "/wp-content/", "/wp-admin/", "/wp-json/"))


def ghep_truong_html(chuoi, ban, doi=lambda u: u):
    """Thay các đoạn của một trường HTML bằng bản Anh; ban = [(đoạn từ doan_json, bản Anh)]. Phần còn lại giữ nguyên từng ký tự."""
    tokens, thay = tach_token(chuoi), {}
    for d, en in ban:
        if d["cach"] == "alt":
            raw = tokens[d["i"]][1]; m = thuoc_tinh(raw, "alt")
            thay[d["i"]] = (d["i"], raw[:m.start(2)] + html.escape(en, quote=True) + raw[m.end(2):])
        else:  # giữ khoảng trắng ở hai mép đoạn (xuống dòng, thụt lề giữa icon và chữ)
            dau = re.match(r"\s*", tokens[d["tu"]][1]).group(0) if tokens[d["tu"]][0] == "chu" else ""
            cuoi = re.search(r"\s*$", tokens[d["den"]][1]).group(0) if tokens[d["den"]][0] == "chu" else ""
            thay[d["tu"]] = (d["den"], dau + sang_html(en, d["tt"], doi) + cuoi)
    ra, i = [], 0
    while i < len(tokens):
        if i in thay:
            ra.append(thay[i][1]); i = thay[i][0] + 1
        else:
            ra.append(tokens[i][1]); i += 1
    return "".join(ra)


def dung_json(viec, dong):
    """Dựng trang tiếng Anh từ nguon.json + bản dịch trong song-ngu.tsv. Trả (trang, lỗi, cảnh báo)."""
    meta = json.loads((viec / "meta.json").read_text(encoding="utf-8"))
    goc = json.loads((viec / "nguon.json").read_text(encoding="utf-8"))[0]
    bc = doc_bo_chuyen(meta["bo_chuyen"])
    ds = doan_json(goc, bc)
    lech = next((i for i, (a, b) in enumerate(zip(ds, dong)) if (a["src"], a["vi"]) != (b["src"], b["vi"])), None)
    if lech is not None or len(ds) != len(dong):
        vt = dong[lech]["id"] if lech is not None else f"đoạn thứ {min(len(ds), len(dong)) + 1}"
        return None, [f"JSON: song-ngu.tsv không còn khớp nguon.json + bộ chuyển {bc['ten']} (lệch từ {vt}) — chạy lại `moi`"], []
    loi, canh, bang, link_viet, id_giu = [], [], doc_lien_ket(), [], []

    def doi(u):
        moi, con_viet = doi_lien_ket(u, bang)
        if con_viet:
            link_viet.append(u)
        return moi

    ra, theo_truong = copy.deepcopy(goc), {}
    for d, r in zip(ds, dong):
        theo_truong.setdefault(d["duong"], []).append((d, r))
    thieu = [r["id"] for r in dong if not (r["en"] or "").strip()]
    if thieu:
        loi.append(f"JSON: còn {len(thieu)} đoạn chưa dịch ({', '.join(thieu[:8])}{'…' if len(thieu) > 8 else ''}) — không dựng được tệp import")
    for duong, cac in theo_truong.items():
        if cac[0][0]["cach"] == "chu":
            en = (cac[0][1]["en"] or "").strip()
            if the_giu_cho(en) or "{br}" in en:
                loi.append(f"{cac[0][1]['id']} THẺ HTML LỆCH: trường chữ thường không được có thẻ giữ chỗ — «{en[:60]}»")
            dat_theo_duong(ra, duong, "" if en == "[BO]" else (en or cac[0][0]["vi"]))
            continue
        ban = []
        for d, r in cac:
            en = (r["en"] or "").strip()
            bo = en == "[BO]"
            en = "" if bo else (en or d["vi"])  # đoạn còn trống giữ chữ Việt (đã báo lỗi ở trên)
            if d["cach"] == "chay" and not bo and the_giu_cho(en) != the_giu_cho(d["vi"]):
                loi.append(f"{r['id']} THẺ HTML LỆCH: nguồn {' '.join(the_giu_cho(d['vi'])) or '(không)'} · bản Anh {' '.join(the_giu_cho(en)) or '(không)'} — giữ đủ thẻ giữ chỗ")
            else:
                ban.append((d, en))
        dat_theo_duong(ra, duong, ghep_truong_html(lay_theo_duong(goc, duong), ban, doi))
    for duong, ten, v in duyet_la(ra.get("acf") or {}, ("acf",)):
        kieu = luat_truong(ten, bc)
        if kieu == "url" and isinstance(v, str) and v:
            dat_theo_duong(ra, duong, doi(v))
        elif kieu == "id" and v not in (None, "", False, []):
            ids = v if isinstance(v, list) else [v]
            moi = [int(bang[f"id:{x}"]) if f"id:{x}" in bang else x for x in ids]
            id_giu += [str(x) for x in ids if f"id:{x}" not in bang]
            dat_theo_duong(ra, duong, moi if isinstance(v, list) else moi[0])
    if not thieu:  # lưới an toàn: dựng lại xong không còn chữ Việt trong các trường đã dịch
        mien = {d["duong"] for d, r in zip(ds, dong) if "#giu-tieng-viet" in (r.get("ghi_chu") or "")}
        for d in doan_json(ra, bc):
            if d["duong"] not in mien and con_tieng_viet(bo_link(d["vi"])):
                loi.append(f"JSON: CÒN CHỮ TIẾNG VIỆT sau khi dựng ở {d['src']}: «{d['vi'][:60]}»")
    if link_viet:
        canh.append(f"JSON: {len(link_viet)} liên kết còn trỏ trang tiếng Việt (chưa có trong lien-ket/lien-ket-vi-en.tsv): "
                    + ", ".join(sorted(set(link_viet))))
    if id_giu:
        canh.append(f"JSON: ID bài giữ nguyên, vẫn trỏ bài tiếng Việt (chưa có trong bảng liên kết): {', '.join(id_giu)}")
    seo_p = viec / "ban-giao" / "seo.json"
    if seo_p.exists():
        seo = json.loads(seo_p.read_text(encoding="utf-8"))
        ra["seo"] = {k: seo.get(k, "") for k in ("meta_title", "meta_description", "slug")}
    return ra, loi, canh


def lenh_moi(args):
    if not args:
        sys.exit("Thiếu nguồn: url hoặc đường dẫn tệp")
    nguon = args[0]
    ten = args[args.index("--ten") + 1] if "--ten" in args else ""
    meta, dong, tho, trang_json = {"nguon": nguon}, [], "", None
    if re.match(r"https?://", nguon):
        tho = lay_url(nguon)
        m, dong = tach_html(tho)
        meta.update(m)
        ten = ten or slugify(urllib.parse.urlparse(nguon).path.strip("/").split("/")[-1] or meta.get("title", ""))
    else:
        p = Path(nguon).expanduser().resolve()
        if not p.exists():
            sys.exit(f"Không thấy tệp: {p}")
        duoi = p.suffix.lower()
        if duoi in (".html", ".htm"):
            tho = p.read_text(encoding="utf-8", errors="replace"); m, dong = tach_html(tho); meta.update(m)
        elif duoi == ".json" or la_json_acf(p):  # JSON dán vào chat có thể bị lưu nhầm đuôi .md/.txt
            dong, m, trang_json, ten_json = tach_json_tep(p, args); meta.update(m); ten = ten or ten_json
        elif duoi == ".docx":
            dong = tach_docx(p)
        elif duoi == ".pdf":
            dong = tach_md(subprocess.run(["pdftotext", "-enc", "UTF-8", str(p), "-"], capture_output=True, text=True).stdout)
        else:
            dong = tach_md(p.read_text(encoding="utf-8", errors="replace"))
        ten = ten or re.sub(r"^\d{4}-\d{2}-\d{2}-", "", slugify(p.stem))
    viec = GOC / "viec" / f"{dt.date.today():%Y-%m-%d}-{ten}"
    so = 2
    while (viec / "song-ngu.tsv").exists():  # đã có việc cùng tên trong ngày → việc mới, không ghi đè bản dịch cũ
        viec = GOC / "viec" / f"{dt.date.today():%Y-%m-%d}-{ten}-{so}"
        so += 1
    viec.mkdir(parents=True, exist_ok=True)
    (viec / "ban-giao").mkdir(exist_ok=True)
    if tho:
        (viec / "nguon.html").write_text(tho, encoding="utf-8")
    if trang_json is not None:
        (viec / "nguon.json").write_text(json.dumps([trang_json], ensure_ascii=False, indent=4) + "\n", encoding="utf-8")
    tm = doc_tm()
    lan_dau = {}
    for i, d in enumerate(dong, 1):
        d["id"] = f"s{i:03d}"
        d["en"] = ""
        d["tm"] = goi_y_tm(d["vi"], tm, loai=d["loai"])
        k = (d["loai"] == "img", chuan(d["vi"]))  # alt ảnh và chữ thường giống nhau vẫn dịch khác nhau
        d["ghi_chu"] = f"trung:{lan_dau[k]}" if k in lan_dau and len(k[1]) > 2 else ""
        lan_dau.setdefault(k, d["id"])
    ghi_tsv(viec / "song-ngu.tsv", dong)
    meta["ngay"] = f"{dt.datetime.now():%Y-%m-%d %H:%M}"
    meta["so_doan"] = len(dong)
    meta["so_tu_vi"] = sum(len(bo_the(d["vi"]).split()) for d in dong)
    en_url = meta.get("hreflang_en")
    if en_url and en_url.rstrip("/") != nguon.rstrip("/"):
        try:
            _, en_dong = tach_html(lay_url(en_url))
            (viec / "tham-khao-ban-en-hien-co.md").write_text(
                f"<!-- Bản tiếng Anh ĐANG có trên web: {en_url} — chỉ để tham khảo, KHÔNG chép -->\n\n"
                + "\n\n".join(f"[{d['loai']}] {d['vi']}" for d in en_dong), encoding="utf-8")
            meta["ban_en_hien_co"] = en_url
        except Exception as e:  # noqa
            meta["ban_en_hien_co_loi"] = str(e)
    (viec / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")
    so_tm = sum(1 for d in dong if d["tm"])
    print(f"VIỆC: {viec}\nĐoạn: {len(dong)} · chữ nguồn: {meta['so_tu_vi']} · khớp bộ nhớ dịch: {so_tm}"
          + (f" · có bản EN cũ: {en_url}" if meta.get("ban_en_hien_co") else ""))
    if len(dong) == 0:
        print("⚠ KHÔNG tách được đoạn nào — trang có thể chặn máy hoặc dựng bằng JavaScript. Lấy chữ bằng Browser pane rồi chạy lại với tệp .md")
        sys.exit(2)


# ---------- từ khoá ----------
def lenh_tukhoa(args):
    viec = Path(args[0]); hat = [a for a in args[1:] if not a.startswith("--")]
    tt = "us,gb,au,ca,sg,ae"
    if "--thi-truong" in args:
        tt = args[args.index("--thi-truong") + 1]; hat = [h for h in hat if h != tt]
    duoi = ["", " requirements", " cost", " for", " vs", " how to", " best", " benefits", " process"]
    kq = {}
    for h in hat:
        for g in tt.split(","):
            for d in duoi:
                q = urllib.parse.quote(h + d)
                url = f"https://suggestqueries.google.com/complete/search?client=firefox&hl=en&gl={g}&q={q}"
                try:
                    goi_y = json.loads(lay_url(url, 10))[1]
                except Exception:
                    continue
                for hang, k in enumerate(goi_y):
                    e = kq.setdefault(k.lower(), {"tu_khoa": k.lower(), "thi_truong": set(), "so_lan": 0, "hang_tot_nhat": 99, "hat": h})
                    e["thi_truong"].add(g); e["so_lan"] += 1; e["hang_tot_nhat"] = min(e["hang_tot_nhat"], hang)
                time.sleep(0.25)
    ds = sorted(kq.values(), key=lambda e: (-len(e["thi_truong"]), -e["so_lan"], e["hang_tot_nhat"]))
    for e in ds:
        e["thi_truong"] = sorted(e["thi_truong"])
    (viec / "tu-khoa.json").write_text(json.dumps(ds, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"{len(ds)} gợi ý từ khoá · thị trường {tt} · ghi {viec/'tu-khoa.json'}")
    print("⚠ Đây là gợi ý tự hoàn thành của Google (người thật đang gõ), KHÔNG có số lượt tìm kiếm.")
    for e in ds[:25]:
        print(f"  {len(e['thi_truong'])} thị trường · {e['so_lan']:>2} lần · {e['tu_khoa']}")


# ---------- tỷ giá ----------
def lenh_tygia(args):
    viec = Path(args[0])
    try:
        d = json.loads(lay_url("https://open.er-api.com/v6/latest/USD", 15))
        r = d["rates"]
        kq = {"nguon": "open.er-api.com (ExchangeRate-API)", "cap_nhat": d.get("time_last_update_utc"),
              "lay_luc": f"{dt.datetime.now():%Y-%m-%d %H:%M}", "vnd_moi_usd": r["VND"],
              "usd_moi_eur": round(1 / r["EUR"], 4), "usd_moi_gbp": round(1 / r["GBP"], 4),
              "usd_moi_aud": round(1 / r["AUD"], 4), "usd_moi_cad": round(1 / r["CAD"], 4),
              "usd_moi_nzd": round(1 / r["NZD"], 4), "usd_moi_sgd": round(1 / r["SGD"], 4)}
    except Exception as e:
        sys.exit(f"LỖI lấy tỷ giá: {e} — thử nguồn Vietcombank bằng tay, ghi ty-gia.json đủ trường vnd_moi_usd")
    (viec / "ty-gia.json").write_text(json.dumps(kq, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"1 USD = {kq['vnd_moi_usd']:,.0f} VND · 1 EUR = {kq['usd_moi_eur']} USD · nguồn {kq['nguon']} · {kq['cap_nhat']}")


# ---------- ghép ----------
def md_inline_html(s):
    s = html.escape(s, quote=False)
    s = re.sub(r"\[([^\]]+)\]\(([^)\s]+)\)", r'<a href="\2">\1</a>', s)
    return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)


def doc_meta(viec):
    p = viec / "meta.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else {}


class KhoiMd:
    """Gom đoạn thành markdown. Bản VI và bản EN dùng chung lớp này nên xuống dòng giống hệt nhau."""

    def __init__(self):
        self.md, self.trong_ul = [], False

    def them(self, loai, chu, anh=""):
        if loai != "li" and self.trong_ul:
            self.trong_ul = False; self.md.append("")
        if re.fullmatch(r"h[1-6]", loai):
            self.md += ["#" * int(loai[1]) + " " + chu, ""]
        elif loai == "li":
            self.md.append(chu if re.match(r"^\d+\.\s", chu) else "- " + chu)  # mục có sẵn số giữ dạng danh sách đánh số
            self.trong_ul = True
        elif loai == "img":
            self.md += [f"![{chu}]({anh})", ""]
        elif loai == "quote":
            self.md += ["> " + chu, ""]
        else:
            self.md += [chu, ""]

    def ra(self):
        return "\n".join(self.md).strip() + "\n"


def lenh_ghep(args):
    viec = Path(args[0]); dong = doc_tsv(viec / "song-ngu.tsv")
    la_json = doc_meta(viec).get("dang") == "acf-json"
    chi_vi = "--chi-vi" in args  # lấp bản VI cho việc cũ, không ghi đè bản EN đã qua cửa 0
    en_md, vi_md = KhoiMd(), KhoiMd()
    h, trong_ul = [], False
    for d in dong:
        en = gon(bo_the((d["en"] or "").strip()))
        if not en or en == "[BO]" or d["loai"] == "tieu-de":  # tiêu đề trang không nằm trong thân bài
            continue
        loai = d["loai"]
        anh = "" if la_json else d["src"]  # việc JSON: cột src là vị trí trường, không phải đường dẫn ảnh
        en_md.them(loai, en, anh)
        vi_md.them(loai, gon(bo_the((d["vi"] or "").strip())), anh)  # bỏ đúng đoạn bản EN bỏ ⇒ hai tệp thẳng hàng từng khối
        if loai != "li" and trong_ul:
            h.append("</ul>"); trong_ul = False
        if re.fullmatch(r"h[1-6]", loai):
            h.append(f"<{loai}>{md_inline_html(en)}</{loai}>")
        elif loai == "li":
            if not trong_ul:
                h.append("<ul>"); trong_ul = True
            h.append(f"  <li>{md_inline_html(en)}</li>")
        elif loai == "img":
            h.append(f'<img src="{html.escape(anh)}" alt="{html.escape(en)}">')
        elif loai == "quote":
            h.append(f"<blockquote>{md_inline_html(en)}</blockquote>")
        else:
            h.append(f"<p>{md_inline_html(en)}</p>")
    if trong_ul:
        h.append("</ul>")
    (viec / "ban-giao").mkdir(exist_ok=True)
    (viec / "ban-giao" / "bai-dich.vi.md").write_text(vi_md.ra(), encoding="utf-8")
    if chi_vi:
        print(f"Chỉ dựng bản VI để đối chiếu → {viec/'ban-giao'/'bai-dich.vi.md'} ({len(vi_md.md)} dòng)")
        return
    (viec / "ban-giao" / "bai-dich.en.md").write_text(en_md.ra(), encoding="utf-8")
    (viec / "ban-giao" / "bai-dich.en.html").write_text("\n".join(h) + "\n", encoding="utf-8")
    print(f"Ghép xong {len(en_md.md)} dòng markdown → {viec/'ban-giao'} (kèm bai-dich.vi.md để đối chiếu)")
    if la_json:
        trang, loi, canh = dung_json(viec, dong)
        if loi:
            print("✗ KHÔNG ghi tệp JSON để import:", *[f"  ✗ {x}" for x in loi], sep="\n")
            sys.exit(1)
        slug = re.sub(r"^\d{4}-\d{2}-\d{2}-", "", viec.name)
        p = viec / "ban-giao" / f"{slug}.en.json"
        p.write_text(json.dumps([trang], ensure_ascii=False, indent=4) + "\n", encoding="utf-8")
        print(f"Tệp JSON để import bằng ACF Page Importer → {p}", *[f"  ⚠ {x}" for x in canh], sep="\n")


# ---------- kiểm (CỬA 0) ----------
# "zero/one" bỏ vì hay là đại từ ("one clear pattern") — gây cảnh báo giả
TU_SO = {w: i for i, w in enumerate("zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen twenty".split()) if i >= 2}
THANG = {m: i for i, m in enumerate("january february march april may june july august september october november december".split(), 1)}
THANG.update({m[:3]: i for m, i in list(THANG.items())})
NHAN_VI = {"tỷ": 1e9, "tỉ": 1e9, "triệu": 1e6, "nghìn": 1e3, "ngàn": 1e3}
NHAN_EN = {"billion": 1e9, "bn": 1e9, "million": 1e6, "m": 1e6, "mn": 1e6, "thousand": 1e3, "k": 1e3}
SO_RE = r"\d+(?:[.,]\d+)*"


def so_vi(tok):
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+(?:,\d+)?", tok):
        return float(tok.replace(".", "").replace(",", "."))
    if re.fullmatch(r"\d+,\d+", tok):
        return float(tok.replace(",", "."))
    if re.fullmatch(r"\d{1,3}(?:,\d{3})+", tok):
        return float(tok.replace(",", ""))
    return float(tok.replace(",", "."))


def so_en(tok):
    if re.fullmatch(r"\d{1,3}(?:,\d{3})+(?:\.\d+)?", tok):
        return float(tok.replace(",", ""))
    if re.fullmatch(r"\d{1,3}(?:\.\d{3})+", tok):
        return float(tok.replace(".", ""))
    return float(tok.replace(",", ""))


def bo_link(s):
    return bo_the(re.sub(r"\]\([^)]*\)", "]", s or ""))


def lay_so(s, ngon_ngu):
    s = bo_link(s)
    kq, vnd = [], []
    if ngon_ngu == "en":
        # tên tháng viết hoa (tiếng Anh luôn viết hoa tên tháng) đứng trước ngày 1–31 hoặc năm 4 chữ số, hoặc sau ngày
        # — tránh "may" (có thể), "EB-5 may", "20 may change"
        ten_thang = "|".join(m.capitalize() for m in sorted(THANG, key=len, reverse=True))
        s = re.sub(r"\b(" + ten_thang + r")\b\.?(?=\s+(?:(?:[1-9]|[12]\d|3[01])(?:st|nd|rd|th)?\b(?!-)|\d{4}\b))", lambda m: f" {THANG[m.group(1).lower()]} ", s)
        s = re.sub(r"(?<![-\w])((?:[1-9]|[12]\d|3[01]))\s+(" + ten_thang + r")\b\.?", lambda m: f"{m.group(1)} {THANG[m.group(2).lower()]} ", s)
        s = re.sub(r"\b(" + "|".join(TU_SO) + r")\b", lambda m: str(TU_SO[m.group(1).lower()]), s, flags=re.I)
        s = re.sub(r"(\d)(st|nd|rd|th)\b", r"\1", s)
    for m in re.finditer(SO_RE + r"(?:\s*(%s)\b)?" % "|".join(sorted((NHAN_VI if ngon_ngu == "vi" else NHAN_EN), key=len, reverse=True)), s, flags=re.I):
        tok, nhan = m.group(0), m.group(1)
        so_tok = re.match(SO_RE, tok).group(0)
        try:
            v = so_vi(so_tok) if ngon_ngu == "vi" else so_en(so_tok)
        except ValueError:
            continue
        if nhan:
            v *= (NHAN_VI if ngon_ngu == "vi" else NHAN_EN)[nhan.lower()]
        sau = s[m.end():m.end() + 12].lower()
        if ngon_ngu == "vi" and re.match(r"\s*(đồng|vnđ|vnd|đ\b)", sau):
            vnd.append(v)
        else:
            kq.append(round(v, 6))
        # số trong ngày 30/09/2026 đã tách riêng từng số nhờ regex (dấu / không nằm trong SO_RE)
    return kq, vnd


def tim_thuat_ngu(vi_thuong, tn):
    """Khớp dài trước; trả về các dòng thuật ngữ xuất hiện trong câu nguồn."""
    con = vi_thuong
    kq = []
    for r in sorted(tn, key=lambda r: -len(r["vi"])):
        k = r["vi"].strip().lower()
        if not k:
            continue
        mau = r"(?<![\wÀ-ỹ])" + re.escape(k) + r"(?![\wÀ-ỹ])"
        if re.search(mau, con, flags=re.I):
            kq.append(r)
            con = re.sub(mau, " ␀ ", con, flags=re.I)
    return kq


def lenh_kiem(args):
    viec = Path(args[0])
    loi, canh = [], []
    dong = doc_tsv(viec / "song-ngu.tsv")
    la_json = doc_meta(viec).get("dang") == "acf-json"
    tn = doc_thuat_ngu()
    ty_gia = json.loads((viec / "ty-gia.json").read_text()) if (viec / "ty-gia.json").exists() else None

    # 1. đủ đoạn
    for d in dong:
        en, ghi = (d["en"] or "").strip(), d["ghi_chu"] or ""
        if not en and d["vi"].strip():
            loi.append(f"{d['id']} THIẾU BẢN DỊCH: {d['vi'][:70]}")
        elif en == "[BO]":
            (canh if re.search(r"#bo:\s*\S", ghi) else loi).append(f"{d['id']} ĐÃ BỎ đoạn — lý do: {ghi or '(KHÔNG ghi lý do #bo:)'}")
    co_dich = [d for d in dong if (d["en"] or "").strip() not in ("", "[BO]")]

    for d in co_dich:
        vi, en, ghi, i = d["vi"], d["en"], d["ghi_chu"] or "", d["id"]
        # 2. số liệu
        sv, vnd = lay_so(vi, "vi")
        se, _ = lay_so(en, "en")
        con_en = list(se)
        for v in sv:
            if v in con_en:
                con_en.remove(v)
            elif "#bo-qua-so" in ghi:
                canh.append(f"{i} số {v:g} không thấy ở bản Anh (đã ghi #bo-qua-so)")
            else:
                loi.append(f"{i} SỐ LỆCH: nguồn có {v:g} mà bản Anh không có — «{vi[:60]}» → «{en[:60]}»")
        for v in vnd:
            if not ty_gia:
                loi.append(f"{i} CÓ TIỀN VND ({v:,.0f}) mà chưa chạy `tygia` để quy đổi"); continue
            usd = v / ty_gia["vnd_moi_usd"]
            hop = [x for x in con_en if abs(x - usd) <= max(0.05 * usd, 1)]
            if hop:
                con_en.remove(hop[0])
            elif "#bo-qua-so" in ghi:
                canh.append(f"{i} VND {v:,.0f} ≈ US${usd:,.0f} không thấy (đã ghi #bo-qua-so)")
            else:
                loi.append(f"{i} QUY ĐỔI SAI/THIẾU: {v:,.0f} VND ≈ US${usd:,.0f} (sai số cho phép 5%) — bản Anh: «{en[:70]}»")
        if con_en and "#bo-qua-so" not in ghi:
            canh.append(f"{i} bản Anh có số không có ở nguồn: {', '.join(f'{x:g}' for x in con_en)}")
        # 3. link và in đậm
        if vi.count("](") != en.count("]("):
            loi.append(f"{i} SỐ LIÊN KẾT LỆCH: nguồn {vi.count('](')} · bản Anh {en.count('](')}")
        for url in ([] if la_json else re.findall(r"\]\((https?://immgroup\.com/(?!en/)[^)]*)\)", en)):  # việc JSON: đổi theo bảng liên kết, báo ở mục 11
            canh.append(f"{i} liên kết còn trỏ trang tiếng Việt: {url} — thay bằng trang /en/ tương ứng nếu có")
        if vi.count("**") != en.count("**"):
            canh.append(f"{i} số chỗ in đậm lệch")
        # 4. còn tiếng Việt / VND
        if con_tieng_viet(bo_link(en)) and "#giu-tieng-viet" not in ghi:
            loi.append(f"{i} CÒN CHỮ TIẾNG VIỆT trong bản Anh: «{en[:70]}»")
        if re.search(r"\b(VND|VNĐ|dong)\b|đồng", bo_link(en)):
            loi.append(f"{i} CÒN ĐƠN VỊ VND trong bản Anh — phải quy đổi sang USD")
        # 5. thuật ngữ
        for r in tim_thuat_ngu(vi.lower(), [t for t in tn if t.get("loai") in ("thuat-ngu", "ten-rieng", "khong-dich")]):
            dich = [x.strip().lower() for x in (r["en"] + "|" + (r.get("bien_the_en") or "")).split("|") if x.strip()]
            if r.get("loai") == "khong-dich":
                dich = [r["vi"].strip().lower()] + dich
            if not any(x in en.lower() for x in dich):
                thong_diep = f"{i} THUẬT NGỮ «{r['vi']}» phải dịch là «{' | '.join(dich)}» — bản Anh: «{en[:70]}»"
                if r.get("trang_thai") == "chot" and "#bo-qua-tn" not in ghi:
                    loi.append(thong_diep)
                else:
                    canh.append(thong_diep.replace("phải dịch", "gợi ý dịch"))
        # 6. từ cấm / ngữ cảnh Việt Nam
        for r in [t for t in tn if t.get("loai") == "cam"]:
            if r["en"] and re.search(r["en"], en, flags=re.I):
                (loi if r.get("trang_thai") == "chot" else canh).append(f"{i} TỪ CẤM «{r['en']}» ({r.get('ghi_chu','')}): «{en[:70]}»")
        if re.search(r"\bVietnam(?:ese)?\b", en, flags=re.I) and "#viet-nam-hop-le" not in ghi:
            canh.append(f"{i} nhắc Việt Nam — xác nhận đây là dữ kiện công ty, không phải ngữ cảnh khách Việt: «{en[:70]}»")
        # 7. độ dài (dấu hiệu bỏ sót)
        if len(vi) > 60:
            ty = len(bo_link(en)) / max(1, len(bo_link(vi)))
            if ty < 0.55 or ty > 1.9:
                canh.append(f"{i} độ dài bất thường (Anh/Việt = {ty:.2f}) — soát bỏ sót hoặc thêm ý")
        if d["loai"] == "img" and len(en) > 125:
            if True:
                canh.append(f"{i} alt ảnh dài {len(en)} ký tự (nên ≤125)")

    for d in dong:
        if d["loai"] == "img" and not (d["en"] or "").strip():
            loi.append(f"{d['id']} ẢNH THIẾU ALT TIẾNG ANH")

    # 8. cấu trúc heading
    h = [int(d["loai"][1]) for d in co_dich if re.fullmatch(r"h[1-6]", d["loai"])]
    seo_p = viec / "ban-giao" / "seo.json"
    phan_doan = seo_p.exists() and json.loads(seo_p.read_text(encoding="utf-8")).get("phan_doan") is True
    if h.count(1) != 1:
        if phan_doan and h.count(1) == 0:
            canh.append("CẤU TRÚC: nguồn là đoạn trích (phan_doan) không có H1 — H1 do bài gốc quyết")
        else:
            loi.append(f"CẤU TRÚC: cần đúng 1 H1, đang có {h.count(1)}")
    for a, b in zip(h, h[1:]):
        if b > a + 1:
            canh.append(f"CẤU TRÚC: heading nhảy từ H{a} xuống H{b}")

    # 9. SEO
    if not seo_p.exists():
        loi.append("SEO: chưa có ban-giao/seo.json")
    else:
        seo = json.loads(seo_p.read_text(encoding="utf-8"))
        for k in ["meta_title", "meta_description", "slug", "tu_khoa_chinh", "thi_truong"]:
            if not seo.get(k):
                loi.append(f"SEO: thiếu trường {k}")
        mt, md_, sl, kw = seo.get("meta_title", ""), seo.get("meta_description", ""), seo.get("slug", ""), (seo.get("tu_khoa_chinh") or "").lower()
        if mt and not 30 <= len(mt) <= 60:
            loi.append(f"SEO: meta title {len(mt)} ký tự (chuẩn 30–60)")
        if md_ and not 120 <= len(md_) <= 160:
            loi.append(f"SEO: meta description {len(md_)} ký tự (chuẩn 120–160)")
        if sl and (not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", sl) or len(sl) > 75):
            loi.append(f"SEO: slug «{sl}» sai dạng (chữ thường không dấu, gạch nối, ≤75 ký tự)")
        if kw:
            h1 = next((d["en"] for d in co_dich if d["loai"] == "h1"), "")
            than = " ".join(d["en"] for d in co_dich if d["loai"] in ("p", "li"))
            dau_bai = " ".join(bo_link(than).split()[:150]).lower()
            if kw not in mt.lower():
                loi.append(f"SEO: từ khoá chính «{kw}» không có trong meta title")
            if kw not in bo_link(h1).lower():
                canh.append(f"SEO: từ khoá chính «{kw}» không có nguyên cụm trong H1")
            if kw not in md_.lower():
                canh.append("SEO: từ khoá chính không có trong meta description")
            if kw not in dau_bai:
                canh.append("SEO: từ khoá chính không có trong 150 chữ đầu bài")
            tu_kw = [w for w in re.split(r"[^a-z0-9]+", kw) if len(w) > 2]
            if sl and not all(w in sl for w in tu_kw):
                canh.append("SEO: slug chưa chứa đủ chữ chính của từ khoá")
            so_lan = bo_link(than).lower().count(kw)
            so_chu = max(1, len(than.split()))
            if so_lan / so_chu > 0.025:
                loi.append(f"SEO: nhồi từ khoá — «{kw}» xuất hiện {so_lan} lần/{so_chu} chữ (>2,5%)")

    # 10. bảng thuật ngữ
    loi += kiem_bang_thuat_ngu(tn)

    # 11. việc JSON: dựng được tệp import (thẻ giữ chỗ đủ, không còn chữ Việt, link/ID còn trỏ bản Việt)
    if la_json:
        _, l, c = dung_json(viec, dong)
        loi += l; canh += c

    bao_cao = [f"# CỬA 0 — {viec.name} — {dt.datetime.now():%Y-%m-%d %H:%M}",
               f"Kết quả: {'ĐỎ' if loi else 'XANH'} · {len(loi)} lỗi chặn · {len(canh)} cảnh báo · {len(co_dich)}/{len(dong)} đoạn đã dịch", "",
               "## LỖI CHẶN (phải sửa hết)", *([f"- {x}" for x in loi] or ["- (không)"]), "",
               "## CẢNH BÁO (đọc từng dòng, sửa hoặc ghi lý do vào báo cáo)", *([f"- {x}" for x in canh] or ["- (không)"])]
    (viec / "kiem-bao-cao.md").write_text("\n".join(bao_cao) + "\n", encoding="utf-8")
    print("\n".join(bao_cao[:2]))
    for x in loi[:30]:
        print("  ✗", x)
    if len(loi) > 30:
        print(f"  … còn {len(loi)-30} lỗi, xem kiem-bao-cao.md")
    sys.exit(1 if loi else 0)


def kiem_bang_thuat_ngu(tn):
    loi, thay = [], {}
    for r in tn:
        if r.get("loai") not in ("thuat-ngu", "ten-rieng", "khong-dich", "cam"):
            loi.append(f"THUẬT NGỮ: loại sai «{r.get('loai')}» ở dòng «{r.get('vi')}»")
        if r.get("trang_thai") not in ("chot", "goi-y"):
            loi.append(f"THUẬT NGỮ: trạng thái sai «{r.get('trang_thai')}» ở dòng «{r.get('vi')}»")
        k = (r.get("loai"), (r.get("vi") or r.get("en") or "").strip().lower())
        if k in thay:
            loi.append(f"THUẬT NGỮ TRÙNG: «{k[1]}» — xoá dòng cũ, giữ dòng mới")
        thay[k] = 1
    return loi


def lenh_kiemtn(_):
    loi = kiem_bang_thuat_ngu(doc_thuat_ngu())
    print(f"Bảng thuật ngữ: {len(doc_thuat_ngu())} dòng · {'ĐỎ' if loi else 'XANH'}")
    for x in loi:
        print("  ✗", x)
    sys.exit(1 if loi else 0)


# ---------- nạp bộ nhớ dịch ----------
def lenh_nap(args):
    viec = Path(args[0])
    dong = [d for d in doc_tsv(viec / "song-ngu.tsv") if (d["en"] or "").strip() not in ("", "[BO]") and d["vi"].strip()]
    meta = json.loads((viec / "meta.json").read_text(encoding="utf-8")) if (viec / "meta.json").exists() else {}
    cu = {khoa_tm(e): e for e in doc_tm()}
    moi = thay = 0
    for d in dong:
        k = khoa_tm(d)
        if k in cu:
            if cu[k]["en"] != d["en"]:
                thay += 1
            del cu[k]  # xoá cũ, thêm mới ở cuối
        else:
            moi += 1
        cu[k] = {"vi": d["vi"], "en": d["en"], "loai": d["loai"], "viec": viec.name,
                 "nguon": meta.get("nguon", ""), "ngay": f"{dt.date.today():%Y-%m-%d}"}
    TM.write_text("".join(json.dumps(e, ensure_ascii=False) + "\n" for e in cu.values()), encoding="utf-8")
    tu = ['<?xml version="1.0" encoding="UTF-8"?>', '<tmx version="1.4"><header creationtool="imm-dich" creationtoolversion="1" datatype="plaintext" segtype="block" adminlang="en" srclang="vi" o-tmf="jsonl"/><body>']
    for e in cu.values():
        tu.append(f'<tu><prop type="viec">{html.escape(e["viec"])}</prop><tuv xml:lang="vi"><seg>{html.escape(e["vi"])}</seg></tuv><tuv xml:lang="en"><seg>{html.escape(e["en"])}</seg></tuv></tu>')
    tu.append("</body></tmx>")
    TMX.write_text("\n".join(tu), encoding="utf-8")
    print(f"Bộ nhớ dịch: +{moi} mới · {thay} cập nhật bản dịch · tổng {len(cu)} · đã xuất TMX")


# ---------- xem / điền ----------
def lenh_xem(args):
    viec = Path(args[0]); dong = doc_tsv(viec / "song-ngu.tsv")
    tu = args[args.index("--tu") + 1] if "--tu" in args else ""
    den = args[args.index("--den") + 1] if "--den" in args else ""
    for d in dong:
        if (tu and d["id"] < tu) or (den and d["id"] > den):
            continue
        ghi = f" {{{d['ghi_chu']}}}" if d["ghi_chu"] else ""
        print(f"[{d['id']}|{d['loai']}]{ghi} {d['vi']}")
        if d["tm"]:
            print(f"    TM {d['tm']}")
        if "--anh" in args:
            print(f"    EN {d['en']}")


def doc_tep_dien(text):
    kq, hien = {}, None
    for l in text.splitlines():
        m = re.match(r"^\[(s\d{3,})(?:\|([^\]]*))?\]\s?(.*)$", l)
        if m:
            hien = m.group(1); kq[hien] = {"en": m.group(3).strip(), "ghi_chu": (m.group(2) or "").strip()}
        elif hien and l.strip() and not l.startswith("    TM "):
            kq[hien]["en"] = (kq[hien]["en"] + " " + l.strip()).strip()
    return kq


def lenh_dien(args):
    viec = Path(args[0]); p = viec / "song-ngu.tsv"; dong = doc_tsv(p)
    theo_id = {d["id"]: d for d in dong}
    dien = {}
    for t in [a for a in args[1:] if not a.startswith("--")]:
        dien.update(doc_tep_dien(Path(t).read_text(encoding="utf-8")))
    if "--trung" in args:  # đoạn "trung:sXXX" dịch y hệt đoạn gốc (đoạn gốc lấy từ tệp điền hoặc bảng)
        for d in dong:
            m = re.match(r"trung:(s\d+)", d["ghi_chu"] or "")
            if m and d["id"] not in dien and (theo_id[m.group(1)]["loai"] == "img") == (d["loai"] == "img"):
                goc = dien.get(m.group(1), {}).get("en") or theo_id[m.group(1)]["en"]
                if goc:
                    dien[d["id"]] = {"en": goc, "ghi_chu": ""}
    if "--tm100" in args:  # đoạn còn trống mà bộ nhớ dịch khớp 100% thì dùng lại nguyên văn
        for d in dong:
            if d["id"] not in dien and not (d["en"] or "").strip() and (d["tm"] or "").startswith("100%: "):
                dien[d["id"]] = {"en": d["tm"][6:], "ghi_chu": ""}
    la = [k for k in dien if k not in theo_id]
    if la:
        sys.exit(f"LỖI: mã đoạn không có trong bảng: {', '.join(la)} — không ghi gì")
    for k, v in dien.items():
        theo_id[k]["en"] = v["en"]
        if v["ghi_chu"]:
            cu = theo_id[k]["ghi_chu"]
            theo_id[k]["ghi_chu"] = (cu + " " + v["ghi_chu"]).strip() if v["ghi_chu"] not in cu else cu
    ghi_tsv(p, dong)
    con = sum(1 for d in dong if not (d["en"] or "").strip())
    print(f"Điền {len(dien)} đoạn · còn trống {con}/{len(dong)}")


# ---------- gắn link bản dịch vào bảng đối chiếu ----------
def slug_trang(viec):
    """Slug cuối của trang gốc: ưu tiên post_slug trong meta.json, rồi tới URL nguồn, cuối cùng là tên thư mục việc."""
    meta = doc_meta(viec)
    s = (meta.get("post_slug") or "").strip("/")
    if s:
        return s.split("/")[-1]
    ng = (meta.get("nguon") or "")
    if ng.startswith("http"):
        duong = urllib.parse.urlparse(ng).path.strip("/")
        if duong:
            return duong.split("/")[-1]
    return re.sub(r"^\d{4}-\d{2}-\d{2}-", "", viec.name)


def lenh_ganlink(args):
    """Ghi link bản dịch trên GitHub vào lien-ket/ban-dich-vi-en.tsv (nguồn sự thật, diff được trong git).
    Bảng Excel của team KHÔNG nằm trong repo: dựng bản mới bằng --xuat, hoặc ghi thêm vào một tệp cụ thể bằng --xlsx."""
    if "--xuat" in args:
        return xuat_bang_excel(Path(args[args.index("--xuat") + 1]))
    thu = "--thu" in args
    xlsx = Path(args[args.index("--xlsx") + 1]) if "--xlsx" in args else None
    sheet = args[args.index("--sheet") + 1] if "--sheet" in args else SHEET_DOI_CHIEU
    repo = args[args.index("--repo") + 1].rstrip("/") if "--repo" in args else REPO_BLOB
    ep_link = args[args.index("--link") + 1] if "--link" in args else ""
    if "--tat-ca" in args:
        viecs = sorted(p for p in (GOC / "viec").iterdir()
                       if p.is_dir() and (p / "ban-giao" / "bai-dich.en.md").exists())
    else:
        viecs, bo_qua = [], False
        for a in args:
            if bo_qua:
                bo_qua = False; continue
            if a.startswith("--"):
                bo_qua = a in ("--xlsx", "--sheet", "--repo", "--link", "--xuat"); continue
            viecs.append(Path(a))
    if not viecs:
        print("✗ Chưa nêu thư mục việc (hoặc dùng --tat-ca)"); sys.exit(2)
    if ep_link and len(viecs) != 1:
        print("✗ --link chỉ dùng cho đúng một việc"); sys.exit(2)
    bang = doc_bang_ban_dich()
    # từ điển slug -> URL trang gốc: gom từ bảng đã có, và từ cột "link" của tệp Excel nếu có nêu --xlsx
    lien = {}
    for d in bang.values():
        lien.setdefault(urllib.parse.urlparse(d["link"]).path.strip("/").split("/")[-1], []).append(d["link"])
    ws = None
    if xlsx:
        ws, c_link, c_vi, c_en = mo_bang_excel(xlsx, sheet)
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, c_link).value
            if isinstance(v, str) and v.startswith("http"):
                khoa = urllib.parse.urlparse(v).path.strip("/").split("/")[-1]
                if v not in lien.setdefault(khoa, []):
                    lien[khoa].append(v)
    xong = loi = 0
    for viec in viecs:
        if not (viec / "ban-giao" / "bai-dich.en.md").exists():
            print(f"  ✗ {viec.name}: chưa có ban-giao/bai-dich.en.md — dịch xong rồi mới gắn link"); loi += 1; continue
        if not (viec / "ban-giao" / "bai-dich.vi.md").exists():
            print(f"  ✗ {viec.name}: thiếu ban-giao/bai-dich.vi.md — chạy `ghep {viec} --chi-vi` trước"); loi += 1; continue
        slug = slug_trang(viec)
        da_nho = next((d["link"] for d in bang.values() if d.get("viec") == viec.name and d.get("link")), "")
        url = ep_link or da_nho
        if not url:
            hit = lien.get(slug, [])
            if len(hit) != 1:
                gap = "chưa biết URL trang gốc" if not hit else f"khớp {len(hit)} URL: {', '.join(hit)}"
                print(f"  ✗ {viec.name}: slug «{slug}» {gap} — chạy lại kèm --link <url trang gốc>"); loi += 1; continue
            url = hit[0]
        vi = f"{repo}/viec/{viec.name}/ban-giao/bai-dich.vi.md"
        en = f"{repo}/viec/{viec.name}/ban-giao/bai-dich.en.md"
        cu = bang.get(url, {})
        if cu and cu.get("viec") and moi_hon(cu["viec"], viec.name):
            print(f"  · {slug:<44} ← {viec.name}: BỎ QUA, URL này đã có bản mới hơn ({cu['viec']})")
            continue
        dau = "≡" if (cu.get("vi"), cu.get("en")) == (vi, en) else ("↻" if cu else "+")
        bang[url] = {"link": url, "vi": vi, "en": en, "viec": viec.name, "ngay": f"{dt.date.today():%Y-%m-%d}"}
        if ws is not None and not thu:
            ghi_dong_excel(ws, c_link, c_vi, c_en, url, vi, en)
        print(f"  {dau} {slug:<44} ← {viec.name}")
        xong += 1
    if not thu and xong:
        ghi_bang_ban_dich(bang)
        if ws is not None:
            ws.parent.save(xlsx)
    dich_den = BANG_BAN_DICH.name + (f" + {xlsx.name}" if xlsx else "")
    print(f"{'(thử) ' if thu else ''}Gắn link: {xong} việc → {dich_den}" + (f" · {loi} việc chưa gắn được" if loi else ""))
    if loi:
        sys.exit(1)


def moi_hon(a, b):
    """Tên thư mục việc mở đầu bằng ngày yyyy-mm-dd nên so chuỗi là so ngày; hậu tố -2 của việc dịch lại cũng lớn hơn."""
    return a != b and a > b


def mo_bang_excel(xlsx, sheet):
    """Mở bảng Excel có sẵn, trả về (worksheet, cột link, cột vi, cột en)."""
    try:
        import openpyxl
    except ImportError:
        print("✗ Thiếu thư viện openpyxl — cài bằng: pip3 install openpyxl"); sys.exit(1)
    if not xlsx.exists():
        print(f"✗ Không thấy bảng Excel: {xlsx}"); sys.exit(1)
    wb = openpyxl.load_workbook(xlsx)
    if sheet not in wb.sheetnames:
        print(f"✗ Không thấy trang tính «{sheet}» trong {xlsx.name} — có: {', '.join(wb.sheetnames)}"); sys.exit(1)
    ws = wb[sheet]
    tieu_de = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}
    for cot in ("link", "vi", "en"):
        if cot not in tieu_de:
            print(f"✗ Trang tính «{sheet}» thiếu cột «{cot}» ở dòng 1"); sys.exit(1)
    return ws, tieu_de["link"], tieu_de["vi"], tieu_de["en"]


def ghi_dong_excel(ws, c_link, c_vi, c_en, url, vi, en):
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, c_link).value
        if isinstance(v, str) and v.rstrip("/") == url.rstrip("/"):
            ws.cell(r, c_vi).value, ws.cell(r, c_en).value = vi, en
            return r
    return 0


def xuat_bang_excel(dich_den):
    """Dựng bảng Excel MỚI từ ban-dich-vi-en.tsv — để gửi team, không đụng tệp làm việc của họ."""
    try:
        import openpyxl
    except ImportError:
        print("✗ Thiếu thư viện openpyxl — cài bằng: pip3 install openpyxl"); sys.exit(1)
    bang = doc_bang_ban_dich()
    if not bang:
        print(f"✗ {BANG_BAN_DICH.name} chưa có dòng nào — chạy `ganlink` trước"); sys.exit(1)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = SHEET_DOI_CHIEU
    ws.append(["link", "vi", "en", "viec", "ngay"])
    for d in sorted(bang.values(), key=lambda x: x["link"]):
        ws.append([d["link"], d["vi"], d["en"], d.get("viec", ""), d.get("ngay", "")])
    for cot, rong in zip("ABCDE", (70, 95, 95, 46, 12)):
        ws.column_dimensions[cot].width = rong
    ws.freeze_panes = "A2"
    dich_den.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dich_den)
    print(f"Xuất {len(bang)} dòng từ {BANG_BAN_DICH.name} → {dich_den}")


def doc_bang_ban_dich():
    if not BANG_BAN_DICH.exists():
        return {}
    with BANG_BAN_DICH.open(encoding="utf-8", newline="") as f:
        return {d["link"]: d for d in csv.DictReader(f, delimiter="\t")}


def ghi_bang_ban_dich(bang):
    BANG_BAN_DICH.parent.mkdir(exist_ok=True)
    with BANG_BAN_DICH.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, ["link", "vi", "en", "viec", "ngay"], delimiter="\t", lineterminator="\n")
        w.writeheader()
        for d in sorted(bang.values(), key=lambda x: x["link"]):
            w.writerow(d)


if __name__ == "__main__":
    lenh = {"moi": lenh_moi, "tukhoa": lenh_tukhoa, "tygia": lenh_tygia, "ghep": lenh_ghep,
            "kiem": lenh_kiem, "nap": lenh_nap, "kiemtn": lenh_kiemtn, "xem": lenh_xem, "dien": lenh_dien,
            "ganlink": lenh_ganlink}
    if len(sys.argv) < 2 or sys.argv[1] not in lenh:
        print(__doc__); sys.exit(2)
    lenh[sys.argv[1]](sys.argv[2:])
